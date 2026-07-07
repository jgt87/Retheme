"""HTML preview of a RETHEMED deck — renders the same file the user
downloads, so preview and PPTX cannot drift apart.

Each output slide is drawn as: the brandkit layout's furniture (background,
static assets, masks, static text) underneath, then the slide's own shape
tree on top at its true positions — text boxes with per-run formatting,
pictures as-is, vector shapes as SVG, tables as HTML grids, charts as
branded markers (HTML cannot render chart graphics; the PPTX has them).
"""

from __future__ import annotations

import base64
import html as html_mod
import json
from pathlib import Path

from lxml import etree
from pptx import Presentation

from ..extractor.layouts import _rpr_props
from ..extractor.svg_convert import shape_to_svg
from ..ooxml import (DEFAULT_CLR_MAP, NS, compose_group_transform, find,
                     findall, parse_xfrm)
from ..units import PX_PER_INCH
from .generator import (_background_css, css_align, css_font, doc_wrapper,
                        layout_furniture_html, safe_hex, version_warning)

_IMAGE_MIME_OK = ("image/png", "image/jpeg", "image/gif", "image/bmp",
                  "image/svg+xml")


def _esc(text: str) -> str:
    return html_mod.escape(text or "")


class _SlideRenderer:
    def __init__(self, slide, layout_rec, manifest, warnings):
        self.slide = slide
        self.part = slide.part
        self.layout = layout_rec or {}
        self.manifest = manifest
        self.warnings = warnings
        self.sw = manifest["slide_geometry"]["width_in"] * 914400
        self.sh = manifest["slide_geometry"]["height_in"] * 914400
        self.theme = {c["theme_slot"]: c["hex"]
                      for c in manifest.get("colors", [])}
        self.fonts = manifest.get("fonts", {})
        self.z = 200  # slide content sits above layout furniture

    # ------------------------------------------------------------ helpers

    def _box(self, xfrm) -> str | None:
        if not xfrm or xfrm.get("x") is None or xfrm.get("cx") is None:
            return None
        style = (f'left:{xfrm["x"] / self.sw * 100:.2f}%;'
                 f'top:{xfrm["y"] / self.sh * 100:.2f}%;'
                 f'width:{max(xfrm["cx"] / self.sw * 100, 0.05):.2f}%;'
                 f'height:{max(xfrm["cy"] / self.sh * 100, 0.05):.2f}%;')
        transforms = []
        if xfrm.get("rot"):
            transforms.append(f'rotate({xfrm["rot"] / 60000:.1f}deg)')
        if xfrm.get("flipH"):
            transforms.append("scaleX(-1)")
        if xfrm.get("flipV"):
            transforms.append("scaleY(-1)")
        if transforms:
            style += f'transform:{" ".join(transforms)};'
        return style

    def _ph_of(self, sp):
        ph = find(sp, ".//p:nvSpPr/p:nvPr/p:ph")
        if ph is None:
            return None
        return {"type": ph.get("type", "obj"), "idx": int(ph.get("idx", 0))}

    def _layout_ph(self, info):
        """Manifest placeholder record on this layout (position + format)."""
        candidates = self.layout.get("placeholders", [])
        for p in candidates:
            if p["type"] == info["type"] and p.get("idx") == info["idx"]:
                return p
        title_like = {"title", "ctrTitle"}
        for p in candidates:
            if info["type"] in title_like and p["type"] in title_like:
                return p
            if p["type"] == info["type"]:
                return p
        return None

    # --------------------------------------------------------------- text

    def _run_html(self, r, fallback: dict) -> str:
        rpr = find(r, "a:rPr")
        props = dict(fallback)
        props.update(_rpr_props(rpr, self.theme, DEFAULT_CLR_MAP, self.fonts))
        style = ""
        font = css_font(props.get("font"))
        if font:
            style += f"font-family:'{font}', sans-serif;"
        if props.get("size_pt"):
            style += f'font-size:{props["size_pt"] * PX_PER_INCH / 72:.1f}px;'
        if props.get("bold"):
            style += "font-weight:bold;"
        if props.get("italic"):
            style += "font-style:italic;"
        color = safe_hex(props.get("color"))
        if color:
            style += f'color:#{color};'
        text = "".join(t.text or "" for t in findall(r, "a:t"))
        return f'<span style="{style}">{_esc(text)}</span>'

    @staticmethod
    def _bullet_marker(ppr, fallback: dict, counters: dict[int, int],
                       lvl: int) -> str:
        """Resolve the paragraph's bullet: explicit pPr buChar/buAutoNum/
        buNone first, else the placeholder level format from the manifest."""
        bu_char = bu_num = None
        explicit_none = False
        if ppr is not None:
            if find(ppr, "a:buNone") is not None:
                explicit_none = True
            el = find(ppr, "a:buChar")
            if el is not None:
                bu_char = el.get("char")
            el = find(ppr, "a:buAutoNum")
            if el is not None:
                bu_num = el.get("type") or "arabicPeriod"
        if not explicit_none and bu_char is None and bu_num is None:
            fb = fallback.get("bullet")
            if isinstance(fb, dict):
                if fb.get("kind") == "char":
                    bu_char = fb.get("char")
                elif fb.get("kind") == "number":
                    bu_num = fb.get("format") or "arabicPeriod"
        if bu_char:
            return _esc(bu_char)
        if bu_num:
            counters[lvl] = counters.get(lvl, 0) + 1
            n = counters[lvl]
            if bu_num.startswith("alpha"):
                return chr(ord("a") + (n - 1) % 26) + "."
            if bu_num.startswith("roman"):
                romans = ["i", "ii", "iii", "iv", "v", "vi", "vii",
                          "viii", "ix", "x"]
                return romans[(n - 1) % 10] + "."
            return f"{n}."
        return ""

    def _text_body_html(self, sp, ph_rec) -> str:
        tx = find(sp, "p:txBody")
        if tx is None:
            return ""
        levels = (ph_rec or {}).get("format_levels", {})
        paras = []
        counters: dict[int, int] = {}
        for p in findall(tx, "a:p"):
            ppr = find(p, "a:pPr")
            lvl = int(ppr.get("lvl", 0)) + 1 if ppr is not None else 1
            fallback = dict(levels.get(str(lvl), levels.get("1", {})))
            fallback.setdefault("font", self.fonts.get("body"))
            fallback.setdefault("size_pt", 18)
            align = None
            if ppr is not None and ppr.get("algn"):
                align = {"l": "left", "ctr": "center", "r": "right",
                         "just": "justify"}.get(ppr.get("algn"))
            if align is None:
                align = css_align(fallback.get("align"))
            chunks = []
            for child in p:
                tag = etree.QName(child).localname
                if tag == "r":
                    chunks.append(self._run_html(child, fallback))
                elif tag == "br":
                    chunks.append("<br>")
            body = "".join(chunks)
            marker = self._bullet_marker(ppr, fallback, counters, lvl) \
                if body.strip() else ""
            if marker:
                body = (f'<span class="bu" style="display:inline-block;'
                        f'min-width:1.1em;">{marker}</span>{body}')
            p_style = "margin:0 0 0.2em 0;line-height:1.2;"
            if align:
                p_style += f"text-align:{align};"
            indent = 0.9 * (lvl - 1)
            if indent:
                p_style += f"padding-left:{indent}em;"
            if marker:
                # hanging indent so wrapped lines align after the marker
                p_style += "text-indent:-1.1em;margin-left:1.1em;"
            paras.append(f'<p style="{p_style}">{body or "&nbsp;"}</p>')
        return "".join(paras)

    # -------------------------------------------------------------- shapes

    def _render_sp(self, sp, group_xfrm) -> str:
        xfrm = parse_xfrm(find(sp, "p:spPr"))
        if xfrm is not None and group_xfrm:
            xfrm = compose_group_transform(xfrm, group_xfrm)
        info = self._ph_of(sp)
        ph_rec = self._layout_ph(info) if info else None

        box = self._box(xfrm)
        if box is None and ph_rec and ph_rec.get("position"):
            pos, size = ph_rec["position"], ph_rec["size"]
            box = (f'left:{pos["x_pct"]*100:.2f}%;top:{pos["y_pct"]*100:.2f}%;'
                   f'width:{size["width_pct"]*100:.2f}%;'
                   f'height:{size["height_pct"]*100:.2f}%;')
        if box is None:
            return ""
        self.z += 1
        out = []

        sp_pr = find(sp, "p:spPr")
        has_geom = sp_pr is not None and (
            find(sp_pr, "a:prstGeom") is not None or
            find(sp_pr, "a:custGeom") is not None)
        has_paint = sp_pr is not None and (
            find(sp_pr, "a:solidFill") is not None or
            find(sp_pr, "a:gradFill") is not None or
            find(sp, "p:style") is not None)
        if has_geom and has_paint and xfrm and xfrm.get("cx"):
            svg, fmt = shape_to_svg(sp, self.theme, DEFAULT_CLR_MAP,
                                    xfrm["cx"], max(xfrm["cy"], 1),
                                    f"s{self.z}")
            if svg and fmt != "raster_fallback":
                svg = svg.replace("<svg ", '<svg style="width:100%;height:100%" ', 1)
                out.append(f'<div style="position:absolute;{box}'
                           f'z-index:{self.z};">{svg}</div>')

        body = self._text_body_html(sp, ph_rec)
        if body.strip():
            self.z += 1
            out.append(f'<div class="rt-text" style="position:absolute;{box}'
                       f'z-index:{self.z};overflow:visible;">{body}</div>')
        return "".join(out)

    def _render_pic(self, pic, group_xfrm) -> str:
        xfrm = parse_xfrm(find(pic, "p:spPr"))
        if xfrm is not None and group_xfrm:
            xfrm = compose_group_transform(xfrm, group_xfrm)
        box = self._box(xfrm)
        blip = find(pic, ".//a:blip")
        if box is None or blip is None:
            return ""
        r_id = blip.get(f"{{{NS['r']}}}embed")
        if not r_id:
            return ""
        try:
            image_part = self.part.related_part(r_id)
        except KeyError:
            return ""
        mime = image_part.content_type
        if mime not in _IMAGE_MIME_OK:
            self.warnings.append(f'an image ({mime}) cannot be shown in the '
                                 f'browser preview (it IS in the PPTX)')
            return ""
        self.z += 1
        b64 = base64.b64encode(image_part.blob).decode()
        return (f'<img src="data:{mime};base64,{b64}" style="position:absolute;'
                f'{box}z-index:{self.z};object-fit:contain;" alt="">')

    def _render_frame(self, frame) -> str:
        # graphicFrame carries its transform as p:xfrm (not a:xfrm in spPr)
        p_xfrm = frame.find(f"{{{NS['p']}}}xfrm")
        xfrm = None
        if p_xfrm is not None:
            off = p_xfrm.find(f"{{{NS['a']}}}off")
            ext = p_xfrm.find(f"{{{NS['a']}}}ext")
            if off is not None and ext is not None:
                xfrm = {"x": int(off.get("x")), "y": int(off.get("y")),
                        "cx": int(ext.get("cx")), "cy": int(ext.get("cy"))}
        box = self._box(xfrm)
        if box is None:
            return ""
        tbl = frame.find(f".//{{{NS['a']}}}tbl")
        if tbl is not None:
            rows_html = []
            for ri, tr in enumerate(tbl.findall(f"{{{NS['a']}}}tr")):
                cells = []
                for tc in tr.findall(f"{{{NS['a']}}}tc"):
                    text = " ".join(t.text or "" for t in
                                    tc.iter(f"{{{NS['a']}}}t")).strip()
                    tag = "th" if ri == 0 else "td"
                    cells.append(f"<{tag}>{_esc(text)}</{tag}>")
                rows_html.append("<tr>" + "".join(cells) + "</tr>")
            self.z += 1
            return (f'<div class="content-table" style="position:absolute;{box}'
                    f'z-index:{self.z};overflow:auto;"><table>'
                    + "".join(rows_html) + "</table></div>")
        # chart / SmartArt / OLE object: branded marker (HTML cannot render
        # these; the downloaded PPTX contains them fully)
        graphic_data = frame.find(f".//{{{NS['a']}}}graphicData")
        uri = (graphic_data.get("uri") or "") if graphic_data is not None else ""
        kind = "Chart" if uri.endswith("/chart") else "Graphic"
        self.z += 1
        return (f'<div class="content-chart" style="position:absolute;{box}'
                f'z-index:{self.z};"><span>{kind}<br>'
                f'<small>fully retained in the PPTX export</small></span></div>')

    def render(self) -> str:
        parts = []

        def visit(container, group_xfrm):
            for child in container:
                tag = etree.QName(child).localname
                if tag == "sp":
                    parts.append(self._render_sp(child, group_xfrm))
                elif tag == "pic":
                    parts.append(self._render_pic(child, group_xfrm))
                elif tag == "graphicFrame":
                    parts.append(self._render_frame(child))
                elif tag == "grpSp":
                    g_xfrm = parse_xfrm(find(child, "p:grpSpPr"))
                    if group_xfrm and g_xfrm:
                        g_xfrm = compose_group_transform(g_xfrm, group_xfrm)
                    visit(child, g_xfrm)

        sp_tree = find(self.slide._element, ".//p:cSld/p:spTree")
        visit(sp_tree, None)
        return "".join(parts)


def render_deck_html(package_dir: str | Path, deck_path: str | Path,
                     title: str = "Deck") -> tuple[str, list[str]]:
    """Render a rethemed deck to HTML using the brandkit package for layout
    furniture and formatting fallbacks. Returns (html, warnings)."""
    package_dir = Path(package_dir)
    manifest = json.loads(
        (package_dir / "template.json").read_text(encoding="utf-8"))
    warnings: list[str] = []
    vw = version_warning(manifest)
    if vw:
        warnings.append(vw)

    layouts_by_name = {l["name"]: l for l in manifest["layouts"]}
    assets_by_id = {a["asset_id"]: a for a in manifest["assets"]}
    geo = manifest["slide_geometry"]
    w_px = round(geo["width_in"] * PX_PER_INCH)
    h_px = round(geo["height_in"] * PX_PER_INCH)

    prs = Presentation(str(deck_path))
    sections = []
    for slide in prs.slides:
        layout_rec = layouts_by_name.get(slide.slide_layout.name, {})
        bg_css = _background_css(layout_rec.get("background", {}), assets_by_id)
        furniture = layout_furniture_html(
            layout_rec, assets_by_id, package_dir, True, warnings,
            (bg_css, w_px, h_px))
        renderer = _SlideRenderer(slide, layout_rec, manifest, warnings)
        content = renderer.render()
        name = html_mod.escape(layout_rec.get("name", "?"))
        sections.append(f'<section class="slide" data-layout="{name}" '
                        f'style="{bg_css}">{furniture}{content}</section>')

    doc = doc_wrapper(manifest, sections, title)
    return doc, list(dict.fromkeys(warnings))
