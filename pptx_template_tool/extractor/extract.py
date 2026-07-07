"""Extraction orchestrator: .pptx/.potx -> template package directory."""

from __future__ import annotations

import shutil
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

# python-pptx only opens packages whose main part is a *presentation*;
# templates (.potx/.potm) and slideshows (.ppsx/.ppsm) are the same format
# with a different declared content type. Patch the type in a temp copy.
_PRESENTATION_CT = ("application/vnd.openxmlformats-officedocument."
                    "presentationml.presentation.main+xml")
_CONVERTIBLE_CTS = (
    "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
    "application/vnd.openxmlformats-officedocument.presentationml.slideshow.main+xml",
    "application/vnd.ms-powerpoint.template.macroEnabled.main+xml",
    "application/vnd.ms-powerpoint.slideshow.macroEnabled.main+xml",
    "application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml",
)


def _as_presentation(path: Path, tmp_dir: Path) -> Path:
    """Return a path python-pptx can open: `path` itself for normal .pptx,
    otherwise a copy in tmp_dir with the content type rewritten."""
    with zipfile.ZipFile(path) as z:
        ct_xml = z.read("[Content_Types].xml").decode("utf-8")
    if not any(ct in ct_xml for ct in _CONVERTIBLE_CTS):
        return path
    for ct in _CONVERTIBLE_CTS:
        ct_xml = ct_xml.replace(ct, _PRESENTATION_CT)
    out = tmp_dir / f"{path.stem}-as-pptx.pptx"
    with zipfile.ZipFile(path) as src, \
            zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = ct_xml.encode("utf-8")
            dst.writestr(item, data)
    return out

from ..ooxml import find, parse_clr_map
from ..units import emu_to_in
from . import classify, grid, layouts, package, preview, theme
from .assets import AssetCollector


def _aspect_ratio(w: int, h: int) -> str:
    from math import gcd
    g = gcd(w, h)
    rw, rh = w // g, h // g
    if (rw, rh) == (4, 3):
        return "4:3"
    if rw / rh == 16 / 9:
        return "16:9"
    return f"{rw}:{rh}"


def _theme_for_master(master) -> object | None:
    """The theme XML root related to a slide master part."""
    from lxml import etree
    part = master.part
    for rel in part.rels.values():
        if rel.reltype == RT.THEME:
            target = rel.target_part
            el = getattr(target, "element", None)
            if el is not None:
                return el
            return etree.fromstring(target.blob)
    return None


def extract_template(pptx_path: str | Path, out_dir: str | Path,
                     progress=None) -> dict:
    """Run full extraction. Returns the manifest dict; writes the package
    (template.json, template.md, assets/, previews/) to out_dir."""
    pptx_path = Path(pptx_path)
    out_dir = Path(out_dir)
    assets_dir = out_dir / "assets"
    assets_dir.mkdir(parents=True, exist_ok=True)

    def report(step: str, frac: float):
        if progress:
            progress(step, frac)

    report("Parsing presentation", 0.05)
    tmp_dir = Path(tempfile.mkdtemp(prefix="pptx-extract-"))
    try:
        parse_path = _as_presentation(pptx_path, tmp_dir)
    except Exception:
        shutil.rmtree(tmp_dir, ignore_errors=True)
        raise
    try:
        return _extract(parse_path, pptx_path, out_dir, assets_dir, report)
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def _extract(parse_path: Path, pptx_path: Path, out_dir: Path,
             assets_dir: Path, report) -> dict:
    prs = Presentation(str(parse_path))
    slide_w, slide_h = prs.slide_width, prs.slide_height
    slide_w_in, slide_h_in = emu_to_in(slide_w), emu_to_in(slide_h)

    # embed the original template so the package alone can later produce
    # native PPTX output on the template's own layouts
    source_dir = out_dir / "source"
    source_dir.mkdir(parents=True, exist_ok=True)
    embedded_source = source_dir / pptx_path.name
    if not embedded_source.exists():
        shutil.copyfile(pptx_path, embedded_source)

    from .. import __version__

    manifest: dict = {
        "extractor_version": __version__,
        "template_name": pptx_path.stem,
        "source_file": pptx_path.name,
        "source_template": f"source/{pptx_path.name}",
        "slide_geometry": {
            "width_in": slide_w_in, "height_in": slide_h_in,
            "aspect_ratio": _aspect_ratio(slide_w, slide_h),
        },
        "colors": [], "fonts": {}, "format_scheme": {},
        "grid": {}, "assets": [], "layouts": [],
        "extraction_notes": [],
    }
    notes = manifest["extraction_notes"]

    collector = AssetCollector(assets_dir, slide_w, slide_h)
    layout_records: list[dict] = []
    layout_elements_for_usage = []
    multi_master = len(prs.slide_masters) > 1
    if multi_master:
        notes.append(f"Template has {len(prs.slide_masters)} slide masters; "
                     "colors/fonts below are from the first master's theme, "
                     "per-layout data is exact.")

    # ------------------------------------------------ theme (first master)
    report("Extracting theme", 0.1)
    first_master = prs.slide_masters[0]
    theme_el = _theme_for_master(first_master)
    clr_map = parse_clr_map(first_master.element)
    theme_colors = theme.parse_theme_colors(theme_el) if theme_el is not None else {}
    manifest["fonts"] = theme.parse_font_scheme(theme_el) if theme_el is not None else {}
    manifest["format_scheme"] = theme.parse_fmt_scheme(
        theme_el, theme_colors, clr_map) if theme_el is not None else {}

    # ------------------------------------------------ masters & layouts
    n_masters = len(prs.slide_masters)
    for mi, master in enumerate(prs.slide_masters):
        m_theme_el = _theme_for_master(master)
        if m_theme_el is None:
            m_theme_el = theme_el
        m_clr_map = parse_clr_map(master.element)
        m_colors = theme.parse_theme_colors(m_theme_el) if m_theme_el is not None else theme_colors
        m_fonts = theme.parse_font_scheme(m_theme_el) if m_theme_el is not None else manifest["fonts"]
        layout_elements_for_usage.append(master.element)

        # master-level assets appear on every layout of this master
        master_tree = find(master.element, ".//p:cSld/p:spTree")
        master_asset_ids = collector.walk_tree(
            master_tree, master.part, "__master__", m_colors, m_clr_map)
        master_texts = layouts.extract_static_texts(
            master_tree, slide_w, slide_h, m_colors, m_clr_map, m_fonts)

        total_layouts = len(master.slide_layouts)
        for li, layout in enumerate(master.slide_layouts):
            report(f"Layout: {layout.name}",
                   0.15 + 0.5 * (mi + li / max(total_layouts, 1)) / n_masters)
            layout_el = layout.element
            layout_elements_for_usage.append(layout_el)

            phs = layouts.extract_placeholders(
                layout_el, master.element, slide_w, slide_h,
                m_colors, m_clr_map, m_fonts)
            bg, bg_blip = layouts.extract_background(
                layout_el, master.element, m_colors, m_clr_map,
                manifest["format_scheme"])

            layout_tree = find(layout_el, ".//p:cSld/p:spTree")
            asset_ids = collector.walk_tree(
                layout_tree, layout.part, layout.name, m_colors, m_clr_map)

            if bg_blip is not None:
                blip, origin = bg_blip
                bg_part = layout.part if origin == "layout" else master.part
                bg_asset = collector.add_background_image(blip, bg_part, layout.name)
                if bg_asset:
                    bg["asset_id"] = bg_asset
                    asset_ids.append(bg_asset)

            # p:sldLayout @showMasterSp="0" hides master shapes on the layout
            shows_master = layout_el.get("showMasterSp", "1") != "0"
            static_assets = (master_asset_ids if shows_master else []) + asset_ids
            static_texts = layouts.extract_static_texts(
                layout_tree, slide_w, slide_h, m_colors, m_clr_map, m_fonts)
            if shows_master:
                static_texts = master_texts + static_texts

            layout_records.append({
                "name": layout.name,
                "master": master.element.get("name") or f"master{mi + 1}",
                "layout_type": layout_el.get("type"),
                "purpose": layouts.infer_purpose(layout_el.get("type"), layout.name),
                "background": bg,
                "placeholders": phs,
                "static_assets": list(dict.fromkeys(static_assets)),
                "static_texts": static_texts,
            })

    # replace the "__master__" owner marker with the actual layout list
    for asset in collector.assets.values():
        if "__master__" in asset["owning_layouts"]:
            asset["owning_layouts"].remove("__master__")
            for rec in layout_records:
                if asset["asset_id"] in rec["static_assets"] and \
                        rec["name"] not in asset["owning_layouts"]:
                    asset["owning_layouts"].append(rec["name"])

    manifest["layouts"] = layout_records

    # ------------------------------------------------ color roles & grid
    report("Inferring color roles", 0.7)
    manifest["colors"] = theme.infer_color_roles(
        theme_colors, layout_elements_for_usage, clr_map)
    manifest["grid"] = grid.derive_grid(layout_records, slide_w, slide_h)

    # ------------------------------------------------ classify assets
    report("Classifying assets", 0.75)
    purposes = {r["name"]: r["purpose"] for r in layout_records}
    total = len(layout_records)
    for asset in collector.assets.values():
        if asset.get("type") is None:
            a_type, confidence = classify.classify_asset(
                asset, slide_w_in, slide_h_in, purposes)
            asset["type"] = a_type
            if confidence == "low":
                asset["needs_review"] = True
        asset["role_notes"] = classify.role_notes(
            asset, slide_w_in, slide_h_in, total)
        asset["owning_layout"] = ("all" if len(asset["owning_layouts"]) >= total
                                  else asset["owning_layouts"][0]
                                  if len(asset["owning_layouts"]) == 1
                                  else asset["owning_layouts"])
        asset.pop("name_hint", None)
        asset.pop("is_image", None)

    field_order = ["asset_id", "type", "source_format", "file_path",
                   "owning_layout", "owning_layouts", "position", "size",
                   "z_order", "crop", "role_notes", "needs_review",
                   "use_bg_fill"]
    manifest["assets"] = [
        {k: a.get(k) for k in field_order}
        for a in sorted(collector.assets.values(),
                        key=lambda a: (a["type"], a["asset_id"]))
    ]
    flagged = sum(1 for a in manifest["assets"] if a.get("needs_review"))
    if flagged:
        notes.append(f"{flagged} asset(s) flagged needs_review "
                     "(low-confidence classification or raster_fallback format).")

    # ------------------------------------------------ previews
    report("Rendering layout previews", 0.85)
    assets_by_id = {a["asset_id"]: a for a in manifest["assets"]}
    mode = preview.render_all_previews(
        parse_path, layout_records, assets_by_id, slide_w_in, slide_h_in,
        out_dir / "previews", notes=notes)
    manifest["preview_mode"] = mode
    if mode == "schematic":
        notes.append("Previews are schematic (LibreOffice not available); "
                     "install LibreOffice + pypdfium2 for rendered previews.")

    # ------------------------------------------------ write package
    report("Writing package", 0.95)
    package.write_package(out_dir, manifest)
    report("Done", 1.0)
    return manifest
