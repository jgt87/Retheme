"""Apply-flow live preview: the draft preview must be rendered server-side
from the user's CURRENT selections (pre-selected defaults included), so
every decision — colors, fonts, layouts, headings, sizes — shows its effect."""

from __future__ import annotations

from pathlib import Path

from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

from pptx_template_tool.webapp import app as webapp
from pptx_template_tool.webapp.app import app

from .test_retheme import _brandkit_with_distinct_theme

client = TestClient(app)


def _deck_with_dominant_color(tmp_path: Path) -> Path:
    from pptx.dml.color import RGBColor

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for i in range(3):
        box = slide.shapes.add_shape(1, Inches(0.5 + i), Inches(1),
                                     Inches(0.8), Inches(0.8))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor.from_string("1E2761")
    deck = tmp_path / "deck.pptx"
    prs.save(str(deck))
    return deck


def test_preview_draft_reflects_choices(tmp_path):
    deck = _deck_with_dominant_color(tmp_path)
    pkg = _brandkit_with_distinct_theme(tmp_path)  # primary 10069F

    r = client.post(
        "/api/apply",
        files={"content_file": ("deck.pptx", deck.read_bytes())},
        data={"package_path": str(pkg)})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out.get("decisions"), "expected confirmable decisions"
    apply_id = out["apply_id"]
    fam = next(d for d in out["decisions"] if d["kind"] == "color_family")
    assert fam["default"] == "rebrand"  # dominant family

    # preview with the pre-selected default: color must be rebranded
    r = client.post(f"/api/apply/{apply_id}/preview-draft",
                    json={"choices": {fam["id"]: fam["default"]}})
    assert r.status_code == 200
    assert "1e2761" not in r.text.lower(), \
        "default 'rebrand' choice must show in the live preview"

    # flipping the choice re-renders with the original color kept
    r2 = client.post(f"/api/apply/{apply_id}/preview-draft",
                     json={"choices": {fam["id"]: "keep"}})
    assert r2.status_code == 200
    assert "1e2761" in r2.text.lower(), \
        "'keep' choice must show the original color in the live preview"

    # repeated identical request is served from cache with the same result
    r3 = client.post(f"/api/apply/{apply_id}/preview-draft",
                     json={"choices": {fam["id"]: "keep"}})
    assert r3.status_code == 200 and r3.text == r2.text


def test_preview_draft_get_renders_defaults(tmp_path):
    """Browsers running a cached older UI issue a GET: it must render the
    preview with every decision's DEFAULT applied (not blank/405)."""
    deck = _deck_with_dominant_color(tmp_path)
    pkg = _brandkit_with_distinct_theme(tmp_path)
    r = client.post(
        "/api/apply",
        files={"content_file": ("deck.pptx", deck.read_bytes())},
        data={"package_path": str(pkg)})
    out = r.json()
    r = client.get(f'/api/apply/{out["apply_id"]}/preview-draft')
    assert r.status_code == 200
    assert "1e2761" not in r.text.lower(), \
        "GET fallback must render the defaults (dominant family rebranded)"


def test_resize_decision_surfaces_and_is_honored(tmp_path):
    """When the deck and brandkit differ in slide size, a 'resize' decision
    appears; choosing 'fit' scales content (surfaced in the warnings)."""
    from pptx.util import Emu

    prs = Presentation()
    prs.slide_width = Emu(12192000)   # 16:9 wide, brandkit is 4:3
    prs.slide_height = Emu(6858000)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(4), Inches(2), Inches(3), Inches(1)) \
        .text_frame.text = "wide content"
    deck = tmp_path / "wide.pptx"
    prs.save(str(deck))
    pkg = _brandkit_with_distinct_theme(tmp_path)

    out = client.post(
        "/api/apply",
        files={"content_file": ("wide.pptx", deck.read_bytes())},
        data={"package_path": str(pkg)}).json()
    apply_id = out["apply_id"]
    resize = next((d for d in out["decisions"] if d["id"] == "resize"), None)
    assert resize is not None, "resize decision must surface on size mismatch"
    assert {o["value"] for o in resize["options"]} == {"keep", "fit"}
    assert resize["default"] == "keep"

    r = client.post(f"/api/apply/{apply_id}/confirm",
                    json={"choices": {"resize": "fit"}})
    assert r.status_code == 200
    assert any("scaled by" in w for w in r.json()["warnings"]), \
        "'fit' choice must scale content and report it"


def test_index_is_not_cacheable():
    r = client.get("/")
    assert r.status_code == 200
    assert r.headers.get("cache-control") == "no-store"


def test_missing_source_flags_approximate_output(tmp_path):
    """A brandkit without an embedded source template forces the rebuild
    path; apply must warn the user prominently that output is approximate."""
    import json

    deck = _deck_with_dominant_color(tmp_path)
    pkg = _brandkit_with_distinct_theme(tmp_path)
    # strip the embedded source -> apply falls back to the rebuild path
    manifest = json.loads((pkg / "template.json").read_text("utf-8"))
    manifest.pop("source_template", None)
    (pkg / "template.json").write_text(json.dumps(manifest), encoding="utf-8")

    out = client.post(
        "/api/apply",
        files={"content_file": ("deck.pptx", deck.read_bytes())},
        data={"package_path": str(pkg)}).json()
    apply_id = out["apply_id"]
    choices = {d["id"]: d["default"] for d in out.get("decisions", [])
               if d.get("default")}
    r = client.post(f"/api/apply/{apply_id}/confirm", json={"choices": choices})
    assert r.status_code == 200, r.text
    warnings = r.json()["warnings"]
    assert any(w.startswith("APPROXIMATE OUTPUT") for w in warnings), warnings


def test_preview_draft_validates_choices(tmp_path):
    deck = _deck_with_dominant_color(tmp_path)
    pkg = _brandkit_with_distinct_theme(tmp_path)
    r = client.post(
        "/api/apply",
        files={"content_file": ("deck.pptx", deck.read_bytes())},
        data={"package_path": str(pkg)})
    apply_id = r.json()["apply_id"]

    r = client.post(f"/api/apply/{apply_id}/preview-draft",
                    json={"choices": {"x": 5}})
    assert r.status_code == 400
    r = client.post("/api/apply/nosuchjob/preview-draft",
                    json={"choices": {}})
    assert r.status_code == 404


def test_evict_oldest_bounds_registry_and_removes_dirs(tmp_path):
    """The in-memory registry stays bounded and evicted jobs' work dirs are
    deleted, so a long-lived server doesn't leak memory or temp files."""
    reg: dict = {}
    dirs = []
    for i in range(5):
        d = tmp_path / f"job{i}"
        d.mkdir()
        (d / "f.txt").write_text("x")
        dirs.append(d)
        reg[f"id{i}"] = {"dir": str(d)}

    webapp._evict_oldest(reg, 3, lambda j: Path(j["dir"]))

    assert list(reg) == ["id2", "id3", "id4"], "oldest two should be evicted"
    assert not dirs[0].exists() and not dirs[1].exists()
    assert dirs[2].exists() and dirs[4].exists()


def test_draft_cache_cleaned_on_confirm(tmp_path):
    """Per-choice draft previews accumulate during confirmation but must be
    removed once the apply is finalized."""
    deck = _deck_with_dominant_color(tmp_path)
    pkg = _brandkit_with_distinct_theme(tmp_path)
    out = client.post(
        "/api/apply",
        files={"content_file": ("deck.pptx", deck.read_bytes())},
        data={"package_path": str(pkg)}).json()
    apply_id = out["apply_id"]
    fam = next(d for d in out["decisions"] if d["kind"] == "color_family")

    for value in ("rebrand", "keep"):
        client.post(f"/api/apply/{apply_id}/preview-draft",
                    json={"choices": {fam["id"]: value}})
    work_dir = Path(webapp.APPLY_JOBS[apply_id]["dir"])
    assert list(work_dir.glob("draft-*.html")), "drafts should exist now"

    r = client.post(f"/api/apply/{apply_id}/confirm",
                    json={"choices": {fam["id"]: "keep"}})
    assert r.status_code == 200
    assert not list(work_dir.glob("draft-*")), "drafts must be cleaned up"
