"""Regression tests for extraction/application nuances discovered on real
corporate templates. Each test builds a synthetic template exercising one
nuance, so a future code change that re-breaks it fails here first.

Covered nuances:
- .potx content type (python-pptx only opens 'presentation' packages)
- gradient backgrounds with stops stored in reverse document order
- true image backgrounds (p:bg > bgPr > blipFill)
- image-filled shapes used as backgrounds (sp with a:blipFill)
- hidden shapes/groups (hidden="1") that PowerPoint never renders
- zero-width/zero-height lines (vertical/horizontal connectors)
- off-canvas authoring guides
- useBgFill background-mask shapes
- raster_fallback assets kept out of generated HTML
- native PPTX generation from the embedded source template
"""

from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from pptx_template_tool.application.generator import generate_html
from pptx_template_tool.application.pptx_generator import generate_pptx
from pptx_template_tool.extractor.extract import extract_template

LAYOUT1 = "ppt/slideLayouts/slideLayout1.xml"
LAYOUT1_RELS = "ppt/slideLayouts/_rels/slideLayout1.xml.rels"
MASTER = "ppt/slideMasters/slideMaster1.xml"

IMAGE_REL = ('<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org'
             '/officeDocument/2006/relationships/image" Target="../media/{img}"/>')


def _png_bytes(color=(30, 90, 160)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (8, 8), color).save(buf, format="PNG")
    return buf.getvalue()


def build_template(tmp_path: Path, *, layout1_bg: str = "",
                   layout1_shapes: str = "", master_shapes: str = "",
                   layout1_rels: str = "", media: dict[str, bytes] | None = None,
                   as_potx: bool = False) -> Path:
    """Build a synthetic template from the python-pptx default, injecting
    raw OOXML into layout 1 / the master."""
    base = tmp_path / "base.pptx"
    Presentation().save(str(base))
    out = tmp_path / ("test.potx" if as_potx else "test.pptx")

    pres_ct = ("application/vnd.openxmlformats-officedocument."
               "presentationml.presentation.main+xml")
    tmpl_ct = ("application/vnd.openxmlformats-officedocument."
               "presentationml.template.main+xml")

    with zipfile.ZipFile(base) as src, \
            zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for name in src.namelist():
            data = src.read(name)
            if name == "[Content_Types].xml":
                if b'Extension="png"' not in data:
                    data = data.replace(
                        b"</Types>",
                        b'<Default Extension="png" ContentType="image/png"/></Types>')
                if as_potx:
                    data = data.replace(pres_ct.encode(), tmpl_ct.encode())
            elif name == LAYOUT1:
                if layout1_bg:
                    data = data.replace(b"<p:spTree>",
                                        layout1_bg.encode() + b"<p:spTree>")
                if layout1_shapes:
                    data = data.replace(b"</p:spTree>",
                                        layout1_shapes.encode() + b"</p:spTree>")
            elif name == LAYOUT1_RELS and layout1_rels:
                data = data.replace(b"</Relationships>",
                                    layout1_rels.encode() + b"</Relationships>")
            elif name == MASTER and master_shapes:
                data = data.replace(b"</p:spTree>",
                                    master_shapes.encode() + b"</p:spTree>")
            dst.writestr(name, data)
        for fname, blob in (media or {}).items():
            dst.writestr(f"ppt/media/{fname}", blob)
    return out


def extract(tmp_path: Path, pptx: Path) -> tuple[dict, Path]:
    out_dir = tmp_path / "pkg"
    manifest = extract_template(pptx, out_dir)
    return manifest, out_dir


def _shape(name: str, x, y, cx, cy, inner: str, attrs: str = "") -> str:
    return (f'<p:sp{attrs}><p:nvSpPr><p:cNvPr id="900" name="{name}"/>'
            f'<p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr>'
            f'<a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
            f'{inner}</p:spPr></p:sp>')


def _line(name: str, x, y, cx, cy) -> str:
    return (f'<p:cxnSp><p:nvCxnSpPr><p:cNvPr id="901" name="{name}"/>'
            f'<p:cNvCxnSpPr/><p:nvPr/></p:nvCxnSpPr><p:spPr>'
            f'<a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
            f'<a:prstGeom prst="line"><a:avLst/></a:prstGeom>'
            f'<a:ln w="6350"><a:solidFill><a:srgbClr val="BFBFBF"/></a:solidFill></a:ln>'
            f'</p:spPr></p:cxnSp>')


# --------------------------------------------------------------- tests

def test_potx_content_type(tmp_path):
    pptx = build_template(tmp_path, as_potx=True)
    manifest, _ = extract(tmp_path, pptx)
    assert len(manifest["layouts"]) > 0


def test_gradient_stops_sorted_and_rendered(tmp_path):
    bg = ('<p:bg><p:bgPr><a:gradFill><a:gsLst>'
          '<a:gs pos="99000"><a:srgbClr val="00A3E0"/></a:gs>'
          '<a:gs pos="1000"><a:srgbClr val="10069F"/></a:gs>'
          '</a:gsLst><a:lin ang="9060000" scaled="0"/></a:gradFill>'
          '<a:effectLst/></p:bgPr></p:bg>')
    pptx = build_template(tmp_path, layout1_bg=bg)
    manifest, pkg = extract(tmp_path, pptx)
    layout1 = manifest["layouts"][0]
    stops = layout1["background"]["stops"]
    assert [s["pos"] for s in stops] == sorted(s["pos"] for s in stops)
    assert stops[0]["hex"] == "10069F"  # dominant stop first

    html, _ = generate_html(pkg, {"slides": [{"layout": layout1["name"],
                                              "title": "x"}]})
    assert "linear-gradient" in html
    assert html.index("#10069F 1%") < html.index("#00A3E0 99%")


def test_background_image(tmp_path):
    bg = ('<p:bg><p:bgPr><a:blipFill><a:blip r:embed="rIdBG1"/>'
          '<a:stretch><a:fillRect/></a:stretch></a:blipFill>'
          '<a:effectLst/></p:bgPr></p:bg>')
    pptx = build_template(
        tmp_path, layout1_bg=bg,
        layout1_rels=IMAGE_REL.format(rid="rIdBG1", img="bg.png"),
        media={"bg.png": _png_bytes()})
    manifest, pkg = extract(tmp_path, pptx)
    layout1 = manifest["layouts"][0]
    assert layout1["background"]["type"] == "image"
    asset_id = layout1["background"]["asset_id"]
    asset = next(a for a in manifest["assets"] if a["asset_id"] == asset_id)
    assert asset["type"] == "background_image"
    assert (pkg / asset["file_path"]).is_file()

    html, _ = generate_html(pkg, {"slides": [{"layout": layout1["name"],
                                              "title": "x"}]})
    assert "data:image/png;base64" in html


def test_image_filled_shape(tmp_path):
    shape = _shape(
        "BG Rect", 0, 0, 9144000, 6858000,
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        '<a:blipFill><a:blip r:embed="rIdBG2"/>'
        '<a:stretch><a:fillRect/></a:stretch></a:blipFill>')
    pptx = build_template(
        tmp_path, layout1_shapes=shape,
        layout1_rels=IMAGE_REL.format(rid="rIdBG2", img="fill.png"),
        media={"fill.png": _png_bytes((160, 60, 30))})
    manifest, _ = extract(tmp_path, pptx)
    full_bleed = [a for a in manifest["assets"]
                  if a["size"] and a["size"]["width_pct"] >= 0.99
                  and a["source_format"] == "raster"]
    assert full_bleed, "image-filled shape was not extracted"
    assert full_bleed[0]["type"] == "background_image"


def test_hidden_shapes_and_groups_skipped(tmp_path):
    hidden_group = (
        '<p:grpSp><p:nvGrpSpPr><p:cNvPr id="910" name="Helper Grid" hidden="1"/>'
        '<p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
        '<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="9144000" cy="6858000"/>'
        '<a:chOff x="0" y="0"/><a:chExt cx="9144000" cy="6858000"/></a:xfrm></p:grpSpPr>'
        + _line("Hidden Line", 0, 914400, 9144000, 0) +
        '</p:grpSp>')
    visible = _line("Visible Line", 0, 1828800, 9144000, 0)
    pptx = build_template(tmp_path, master_shapes=hidden_group + visible)
    manifest, _ = extract(tmp_path, pptx)
    names = " ".join(a["asset_id"] for a in manifest["assets"])
    assert "visible-line" in names
    assert "hidden-line" not in names
    assert "helper" not in names


def test_zero_extent_lines_kept_and_rendered(tmp_path):
    shapes = (_line("Vert Line", 914400, 0, 0, 6858000) +
              _line("Horiz Line", 0, 914400, 9144000, 0) +
              _line("Offcanvas Tick", -914400, 914400, 91440, 0))
    pptx = build_template(tmp_path, layout1_shapes=shapes)
    manifest, pkg = extract(tmp_path, pptx)
    ids = [a["asset_id"] for a in manifest["assets"]]
    assert "vert-line" in ids
    assert "horiz-line" in ids
    assert "offcanvas-tick" not in ids

    layout1 = manifest["layouts"][0]
    html, _ = generate_html(pkg, {"slides": [{"layout": layout1["name"],
                                              "title": "x"}]})
    assert "min-width:1px;min-height:1px" in html
    assert 'data-asset="vert-line"' in html


def test_use_bg_fill_mask(tmp_path):
    mask = _shape("Hide Logo", 0, 6400800, 914400, 457200,
                  '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
                  '<a:ln><a:noFill/></a:ln>', attrs=' useBgFill="1"')
    pptx = build_template(tmp_path, layout1_shapes=mask)
    manifest, pkg = extract(tmp_path, pptx)
    masks = [a for a in manifest["assets"] if a.get("use_bg_fill")]
    assert masks and masks[0]["source_format"] == "background_mask"
    assert masks[0]["file_path"] is None

    layout1 = manifest["layouts"][0]
    html, _ = generate_html(pkg, {"slides": [{"layout": layout1["name"],
                                              "title": "x"}]})
    assert "bg-mask" in html


def test_raster_fallback_not_in_html(tmp_path):
    weird = _shape("Weird Gear", 914400, 914400, 914400, 914400,
                   '<a:prstGeom prst="gear6"><a:avLst/></a:prstGeom>'
                   '<a:solidFill><a:srgbClr val="FF0000"/></a:solidFill>')
    pptx = build_template(tmp_path, layout1_shapes=weird)
    manifest, pkg = extract(tmp_path, pptx)
    fallback = [a for a in manifest["assets"]
                if a["source_format"] == "raster_fallback"]
    assert fallback and fallback[0]["needs_review"]

    layout1 = manifest["layouts"][0]
    html, warnings = generate_html(pkg, {"slides": [{"layout": layout1["name"],
                                                     "title": "x"}]})
    assert f'data-asset="{fallback[0]["asset_id"]}"' not in html
    assert any("skipped" in w for w in warnings)


def test_native_pptx_generation(tmp_path):
    pptx = build_template(tmp_path, as_potx=True)
    manifest, pkg = extract(tmp_path, pptx)
    assert manifest.get("source_template")  # source embedded in package
    content = {"slides": [
        {"layout": "title", "title": "T", "subtitle": "S"},
        {"layout": "content", "title": "C", "bullets": ["a", {"text": "b",
                                                              "children": ["c"]}]},
        {"layout": "definitely_not_a_layout", "title": "X"},
    ]}
    out = tmp_path / "deck.pptx"
    warnings = generate_pptx(pkg, content, out)
    prs = Presentation(str(out))
    assert len(prs.slides) == 3
    texts = [sh.text_frame.text for sh in prs.slides[0].shapes
             if sh.has_text_frame]
    assert any("T" == t for t in texts)
    assert any("no layout matches" in w for w in warnings)


def test_corrections_regenerate_markdown(tmp_path):
    """Review corrections must keep template.md in sync with template.json."""
    from pptx_template_tool.extractor.package import write_package
    pptx = build_template(tmp_path)
    manifest, pkg = extract(tmp_path, pptx)
    manifest["colors"][0]["role"] = "primary"
    write_package(pkg, manifest)
    md = (pkg / "template.md").read_text(encoding="utf-8")
    json_text = (pkg / "template.json").read_text(encoding="utf-8")
    assert "**primary**" in md
    assert json.loads(json_text)["colors"][0]["role"] == "primary"


def test_classification_confidence(tmp_path):
    """Lines are decorative with high confidence regardless of layout count;
    shapes the template author named 'logo' classify as logo unflagged."""
    logo_shape = _shape(
        "Company Logo", 8001000, 6300000, 548640, 182880,
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        '<a:solidFill><a:srgbClr val="10069F"/></a:solidFill>')
    line = _line("Divider", 0, 3000000, 9144000, 0)
    pptx = build_template(tmp_path, layout1_shapes=logo_shape + line)
    manifest, _ = extract(tmp_path, pptx)
    by_id = {a["asset_id"]: a for a in manifest["assets"]}
    logo = next(a for i, a in by_id.items() if "logo" in i)
    divider = next(a for i, a in by_id.items() if "divider" in i)
    assert logo["type"] == "logo" and not logo["needs_review"]
    assert divider["type"] == "decorative_shape" and not divider["needs_review"]


def test_group_flip_mirrors_children(tmp_path):
    """A grpSp with flipH mirrors its children's positions (PowerPoint
    renders the group mirrored; extraction must match)."""
    group = (
        '<p:grpSp><p:nvGrpSpPr><p:cNvPr id="920" name="Deco Group"/>'
        '<p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
        '<p:grpSpPr><a:xfrm flipH="1"><a:off x="0" y="0"/>'
        '<a:ext cx="9144000" cy="6858000"/><a:chOff x="0" y="0"/>'
        '<a:chExt cx="9144000" cy="6858000"/></a:xfrm></p:grpSpPr>'
        + _line("Right Edge Line", 8229600, 0, 0, 6858000) +  # x = 9in
        '</p:grpSp>')
    pptx = build_template(tmp_path, layout1_shapes=group)
    manifest, _ = extract(tmp_path, pptx)
    line = next(a for a in manifest["assets"] if "right-edge" in a["asset_id"])
    # child sat at x=9in inside a 10in-wide flipped group -> lands at x=1in
    assert abs(line["position"]["x_in"] - 1.0) < 0.01


def _pic(name: str, x, y, cx, cy, rid: str) -> str:
    """A true p:pic element (p:blipFill, not a:blipFill)."""
    return (f'<p:pic><p:nvPicPr><p:cNvPr id="930" name="{name}"/>'
            f'<p:cNvPicPr/><p:nvPr/></p:nvPicPr>'
            f'<p:blipFill><a:blip r:embed="{rid}"/>'
            f'<a:stretch><a:fillRect/></a:stretch></p:blipFill>'
            f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/>'
            f'<a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
            f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>')


def test_true_picture_elements_extracted(tmp_path):
    """p:pic uses p:blipFill (PresentationML ns); the extractor must find
    the a:blip inside it — regression for pictures silently dropped."""
    pic = _pic("Corner Logo", 8001000, 457200, 548640, 274320, "rIdP1")
    pptx = build_template(
        tmp_path, layout1_shapes=pic,
        layout1_rels=IMAGE_REL.format(rid="rIdP1", img="corner.png"),
        media={"corner.png": _png_bytes((10, 200, 10))})
    manifest, pkg = extract(tmp_path, pptx)
    pics = [a for a in manifest["assets"]
            if a["source_format"] == "raster" and a["position"]]
    assert pics, "true p:pic element was not extracted"
    pic_asset = pics[0]
    assert abs(pic_asset["position"]["x_in"] - 8.75) < 0.01
    assert (pkg / pic_asset["file_path"]).is_file()


def test_audit_finds_missing_and_passes_clean(tmp_path):
    """The audit must report zero missing elements on a clean extraction."""
    from pptx_template_tool.extractor.audit import audit_package
    shapes = (_line("Deco Line", 0, 3000000, 9144000, 0) +
              _pic("Photo", 914400, 914400, 1828800, 914400, "rIdP2"))
    pptx = build_template(
        tmp_path, layout1_shapes=shapes,
        layout1_rels=IMAGE_REL.format(rid="rIdP2", img="p.png"),
        media={"p.png": _png_bytes()})
    manifest, pkg = extract(tmp_path, pptx)
    report = audit_package(pptx, pkg)
    assert report["total_missing"] == 0, report
