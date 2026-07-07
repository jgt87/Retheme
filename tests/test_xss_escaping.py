"""Security regression: font names, colors and alignments taken from an
untrusted deck/template must never reach the HTML preview or exported HTML
in a way that can break out of a CSS context and inject markup/script."""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt

from pptx_template_tool.application.generator import (css_align, css_font,
                                                      safe_hex)
from pptx_template_tool.application.preview_render import render_deck_html
from pptx_template_tool.application.retheme import retheme_pptx
from pptx_template_tool.ooxml import NS, resolve_color

from .test_retheme import _brandkit_with_distinct_theme

A = NS["a"]

SCRIPT_PAYLOAD = 'Arial"><script>alert(1)</script>'
COLOR_PAYLOAD = '"><svg onload=alert(2)>'


# ---- unit: the CSS-value sanitizers --------------------------------------

def test_safe_hex_accepts_only_6_hex():
    assert safe_hex("10069F") == "10069F"
    assert safe_hex("abcdef") == "abcdef"
    assert safe_hex("FFF") is None            # too short
    assert safe_hex("1234567") is None        # too long
    assert safe_hex('"><script>') is None
    assert safe_hex("12345G") is None         # non-hex digit
    assert safe_hex(None) is None


def test_css_font_strips_dangerous_chars():
    assert css_font("Segoe UI") == "Segoe UI"
    assert css_font("Helvetica-Neue") == "Helvetica-Neue"
    # quotes, angle brackets, semicolons, braces, slashes all removed
    assert '"' not in css_font(SCRIPT_PAYLOAD)
    assert "<" not in css_font(SCRIPT_PAYLOAD)
    assert ">" not in css_font(SCRIPT_PAYLOAD)
    assert "/" not in css_font(SCRIPT_PAYLOAD)
    assert css_font("';}</style>") is None or "<" not in css_font("';}</style>")
    assert css_font(None) is None


def test_css_align_whitelist():
    assert css_align("center") == "center"
    assert css_align("justify") == "justify"
    assert css_align("center;}</style><script>") is None
    assert css_align(None) is None


# ---- unit: color resolution rejects non-hex at the source ----------------

def test_resolve_color_rejects_non_hex():
    ok = etree.fromstring(
        f'<a:solidFill xmlns:a="{A}"><a:srgbClr val="ABCDEF"/></a:solidFill>')
    assert resolve_color(ok, {}, {})["hex"] == "ABCDEF"

    bad = etree.fromstring(
        f'<a:solidFill xmlns:a="{A}"><a:srgbClr/></a:solidFill>')
    bad[0].set("val", COLOR_PAYLOAD)  # set via API: raw markup, no XML parse
    assert resolve_color(bad, {}, {}) is None

    # a crafted theme slot value is rejected too (schemeClr -> theme hex)
    scheme = etree.fromstring(
        f'<a:solidFill xmlns:a="{A}"><a:schemeClr val="accent1"/></a:solidFill>')
    assert resolve_color(scheme, {"accent1": COLOR_PAYLOAD}, {}) is None
    assert resolve_color(scheme, {"accent1": "112233"}, {})["hex"] == "112233"


# ---- integration: a malicious deck cannot inject into the preview --------

def test_malicious_font_and_color_do_not_inject(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "SAFEWORD"
    run.font.size = Pt(24)
    run.font.name = SCRIPT_PAYLOAD                      # crafted typeface

    # a filled shape whose fill color is a crafted srgbClr value
    from pptx.dml.color import RGBColor
    sq = slide.shapes.add_shape(1, Inches(1), Inches(3), Inches(2), Inches(2))
    sq.fill.solid()
    sq.fill.fore_color.rgb = RGBColor.from_string("112233")  # creates srgbClr
    srgb = sq._element.spPr.find(f".//{{{A}}}srgbClr")
    srgb.set("val", COLOR_PAYLOAD)

    deck = tmp_path / "malicious.pptx"
    prs.save(str(deck))

    pkg = _brandkit_with_distinct_theme(tmp_path)
    out = tmp_path / "branded.pptx"
    retheme_pptx(pkg, deck, out)
    html, _ = render_deck_html(pkg, out, "deck")

    lower = html.lower()
    # the payloads rely on parentheses/brackets that the sanitizers strip
    # (font) or on a value the source rejects entirely (color), so the
    # executable signatures must not survive anywhere in the output
    assert "<script>alert" not in lower, "script tag injected via font name"
    assert "alert(1)" not in lower, "font-name payload survived"
    assert "<svg onload" not in lower, "svg tag injected via fill color"
    assert "alert(2)" not in lower, "fill-color payload survived"
    # the legitimate text content must still be present
    assert "SAFEWORD" in html


def test_malicious_theme_font_does_not_inject_style_block(tmp_path):
    """A crafted theme font would land in the document's <style> block; it
    must be neutralized so it cannot carry </style><script>."""
    import zipfile as zf

    base = _brandkit_with_distinct_theme(tmp_path)  # has a template.json
    manifest_path = base / "template.json"
    import json
    manifest = json.loads(manifest_path.read_text("utf-8"))
    manifest.setdefault("fonts", {})["body"] = 'X";}</style><script>alert(3)</script>'
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    deck = tmp_path / "d.pptx"
    prs.save(str(deck))
    out = tmp_path / "o.pptx"
    retheme_pptx(base, deck, out)
    html, _ = render_deck_html(base, out, "t")

    assert "</style><script>alert(3)" not in html.lower()
    assert "alert(3)" not in html.lower()
