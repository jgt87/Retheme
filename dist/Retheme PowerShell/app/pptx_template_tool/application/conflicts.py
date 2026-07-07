"""Detect brandkit-vs-content-deck decisions that benefit from user
confirmation, and describe them as choice questions the UI can render.

Each decision: {"id", "kind", "label", "detail", "options": [{"value",
"label"}...], "default", "swatches"?: [hex...]}. `retheme_pptx(choices=...)`
consumes the answers keyed by decision id.
"""

from __future__ import annotations

import json
import tempfile
from pathlib import Path

from pptx import Presentation

from ..extractor.layouts import infer_purpose
from .content_extract import _PURPOSE_TO_TAG
from .generator import LayoutNotFoundError, _find_layout
from .retheme import (_harmonization_map, _theme_colors_all, _theme_colors_of,
                      _theme_fonts_all, _theme_fonts_of,
                      collect_hardcoded_fonts, color_clusters,
                      deck_has_textbox_headings)


def analyze_conflicts(package_dir: str | Path,
                      deck_path: str | Path) -> list[dict]:
    """Returns the list of user-confirmable decisions (possibly empty)."""
    from ..extractor.extract import _as_presentation
    import shutil

    package_dir = Path(package_dir)
    manifest = json.loads(
        (package_dir / "template.json").read_text(encoding="utf-8"))
    source_rel = manifest.get("source_template")
    source_path = package_dir / source_rel if source_rel else None
    if not source_path or not source_path.is_file():
        return []  # old package: rebuild path has no confirmable choices

    tmp = Path(tempfile.mkdtemp(prefix="conflicts-"))
    try:
        src = Presentation(str(_as_presentation(Path(deck_path), tmp)))
        dest = Presentation(str(_as_presentation(source_path, tmp)))
        decisions: list[dict] = []

        dest_theme = _theme_colors_of(dest)
        # aggregate across ALL source masters so colors bound to any master's
        # theme are treated as theme-mapped (auto-rebranded), matching the
        # retheme engine — not surfaced as a keep/rebrand family decision
        src_themes = _theme_colors_all(src)
        primary = next((c["hex"] for c in manifest.get("colors", [])
                        if c.get("role") == "primary"), None) \
            or dest_theme.get("accent1", "000000")

        # ---- color families -------------------------------------------
        theme_mapped = {st[s] for st in src_themes for s in st
                        if s in dest_theme and st[s]}
        protected = set(dest_theme.values()) | theme_mapped
        for cluster in color_clusters(src, protected):
            swatches = list(cluster["colors"])[:6]
            role = "dominant color family" if cluster["dominant"] else \
                "color family"
            # old->new mapping so the UI can preview the rebrand live
            mapping = _harmonization_map([cluster], primary, {cluster["id"]})
            decisions.append({
                "id": cluster["id"], "kind": "color_family",
                "label": f'Deck {role} '
                         f'(used {cluster["weight"]}x)',
                "detail": "Rebranding preserves light/dark variations. "
                          "Keep it if these colors carry meaning "
                          "(status, severity, categories).",
                "swatches": swatches,
                "mapping": mapping,
                "options": [
                    {"value": "rebrand",
                     "label": f"Recolor to brand primary #{primary}"},
                    {"value": "keep", "label": "Keep original colors"},
                ],
                "default": "rebrand" if cluster["dominant"] else "keep",
            })

        # ---- foreign fonts --------------------------------------------
        dest_fonts = _theme_fonts_of(dest)
        src_fonts_list = _theme_fonts_all(src)
        exclude = {f for sf in src_fonts_list for f in sf.values()} \
            | set(dest_fonts.values())
        for face, count in sorted(collect_hardcoded_fonts(src, exclude).items(),
                                  key=lambda kv: -kv[1]):
            decisions.append({
                "id": f"font:{face}", "kind": "font",
                "label": f'Font "{face}" (used {count}x)',
                "detail": "This font is not part of the brandkit.",
                "source": face,
                "target": dest_fonts.get("minor", ""),
                "options": [
                    {"value": "map",
                     "label": f'Replace with brand font '
                              f'"{dest_fonts.get("minor", "?")}"'},
                    {"value": "keep", "label": "Keep this font"},
                ],
                "default": "map",
            })

        # ---- layouts without an exact match ---------------------------
        dest_layout_names = []
        for master in dest.slide_masters:
            for layout in master.slide_layouts:
                if layout.name not in dest_layout_names:
                    dest_layout_names.append(layout.name)
        seen: set[str] = set()
        for slide in src.slides:
            lname = slide.slide_layout.name
            if lname in seen or lname in dest_layout_names:
                continue
            seen.add(lname)
            purpose = infer_purpose(slide.slide_layout.element.get("type"),
                                    lname)
            tag = _PURPOSE_TO_TAG.get(purpose, "content")
            try:
                proposed = _find_layout(manifest, tag)["name"]
            except LayoutNotFoundError:
                proposed = dest_layout_names[0]
            decisions.append({
                "id": f"layout:{lname}", "kind": "layout",
                "label": f'Deck layout "{lname}" has no brandkit match',
                "detail": "Choose which brandkit layout these slides use.",
                "options": [{"value": n, "label": n}
                            for n in dest_layout_names],
                "default": proposed,
            })

        # ---- heading text in plain text boxes --------------------------
        if deck_has_textbox_headings(src):
            decisions.append({
                "id": "heading_style", "kind": "toggle",
                "label": "Headings in text boxes",
                "detail": "This deck styles headings as plain text boxes "
                          "(not title placeholders). Convert moves each "
                          "slide's main heading to the layout's title "
                          "position and relocates secondary text in the "
                          "title zone to the brandkit subtitle slot. "
                          "Restyle in place changes only font, size and "
                          "color — nothing moves, so nothing can collide. "
                          "Either way, text sitting on a colored shape "
                          "stays with its shape and keeps its own color "
                          "(the shape's fill follows the color decisions "
                          "above).",
                "options": [
                    {"value": "convert",
                     "label": "Convert to brandkit titles (font, size, "
                              "color + position; secondary text moves to "
                              "the subtitle slot)"},
                    {"value": "adopt",
                     "label": "Restyle in place (font, size, color only "
                              "— no movement)"},
                    {"value": "keep", "label": "Keep the deck's heading style"},
                ],
                "default": "convert",
            })

        # ---- slide-size mismatch ---------------------------------------
        if (src.slide_width, src.slide_height) != \
                (dest.slide_width, dest.slide_height):
            decisions.append({
                "id": "resize", "kind": "toggle",
                "label": "Slide sizes differ",
                "detail": "The content deck and brandkit use different slide "
                          "dimensions. Keep positions leaves every shape at "
                          "its original coordinates (you may need to nudge "
                          "content in PowerPoint). Scale to fit uniformly "
                          "resizes content onto the brand canvas without "
                          "stretching or distorting shapes.",
                "options": [
                    {"value": "keep", "label": "Keep original positions"},
                    {"value": "fit",
                     "label": "Scale content to fit the brand canvas"},
                ],
                "default": "keep",
            })

        # ---- title sizing ----------------------------------------------
        decisions.append({
            "id": "title_sizes", "kind": "toggle",
            "label": "Title font sizes",
            "detail": "Adopt the brandkit layout's title sizes, or keep the "
                      "sizes authored in the content deck.",
            "options": [
                {"value": "adopt", "label": "Use brandkit title sizes"},
                {"value": "keep", "label": "Keep the deck's title sizes"},
            ],
            "default": "adopt",
        })
        return decisions
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
