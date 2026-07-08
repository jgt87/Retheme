"""Retheme: apply a brandkit to a user's content deck while RETAINING all
slide content — pictures, charts, tables, shapes, icons, embedded objects.

Unlike the rebuild path (generator/pptx_generator, used for LLM content
JSON), this copies each slide of the content deck wholesale into the
brandkit's embedded source template:

- every shape/picture/graphicFrame is deep-copied, with its related parts
  (images, chart XML, chart data workbooks, SmartArt) cloned into the
  destination package;
- the slide's layout reference is remapped to the matching brandkit layout,
  so theme-bound colors (schemeClr) and fonts re-resolve against the brand
  theme automatically;
- charts that hard-code colors equal to the source deck's theme accents get
  those colors rewritten to the brandkit's corresponding accents;
- pictures are kept byte-for-byte as-is.
"""

from __future__ import annotations

import colorsys
import copy
import hashlib
import json
import re
import shutil
import tempfile
from collections import Counter
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI

from ..extractor.layouts import infer_purpose
from ..ooxml import NS
from .content_extract import _PURPOSE_TO_TAG
from .generator import LayoutNotFoundError, _find_layout
from .pptx_generator import _delete_all_slides

R_NS = NS["r"]
A_NS = NS["a"]
P_NS = NS["p"]

_THEME_SLOTS = ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3",
                "accent4", "accent5", "accent6")
_CHART_CT_MARKERS = ("chart+xml", "chartstyle+xml", "chartcolorstyle+xml")


def _theme_root_of(prs):
    for rel in prs.slide_masters[0].part.rels.values():
        if rel.reltype == RT.THEME:
            return etree.fromstring(rel.target_part.blob)
    return None


def _theme_roots_all(prs) -> list:
    """Theme root of every slide master (a template may carry several, each
    with its own color/font scheme)."""
    roots = []
    for master in prs.slide_masters:
        for rel in master.part.rels.values():
            if rel.reltype == RT.THEME:
                roots.append(etree.fromstring(rel.target_part.blob))
                break
    return roots


def _colors_from_root(root) -> dict[str, str]:
    out: dict[str, str] = {}
    scheme = root.find(f".//{{{A_NS}}}clrScheme") if root is not None else None
    if scheme is None:
        return out
    for slot in _THEME_SLOTS:
        el = scheme.find(f"{{{A_NS}}}{slot}")
        if el is None:
            continue
        srgb = el.find(f"{{{A_NS}}}srgbClr")
        sysc = el.find(f"{{{A_NS}}}sysClr")
        if srgb is not None:
            out[slot] = srgb.get("val", "").upper()
        elif sysc is not None:
            out[slot] = (sysc.get("lastClr") or "").upper()
    return out


def _fonts_from_root(root) -> dict[str, str]:
    out: dict[str, str] = {}
    scheme = root.find(f".//{{{A_NS}}}fontScheme") if root is not None else None
    if scheme is None:
        return out
    for kind in ("major", "minor"):
        latin = scheme.find(f"{{{A_NS}}}{kind}Font/{{{A_NS}}}latin")
        if latin is not None and latin.get("typeface"):
            out[kind] = latin.get("typeface")
    return out


def _theme_colors_of(prs) -> dict[str, str]:
    """{slot: hex} for the first master's theme of a presentation."""
    return _colors_from_root(_theme_root_of(prs))


def _theme_fonts_of(prs) -> dict[str, str]:
    """{"major": typeface, "minor": typeface} of the first master's theme."""
    return _fonts_from_root(_theme_root_of(prs))


def _theme_colors_all(prs) -> list[dict]:
    """One {slot: hex} per slide master (multi-master templates)."""
    return [_colors_from_root(r) for r in _theme_roots_all(prs)]


def _theme_fonts_all(prs) -> list[dict]:
    """One {major/minor: typeface} per slide master."""
    return [_fonts_from_root(r) for r in _theme_roots_all(prs)]


def _source_theme_color_map(src_themes: list[dict],
                            dest_theme: dict) -> dict[str, str]:
    """Map every source master theme's slot color to the brand's color for the
    same slot, so a hard-coded color matching ANY source master's theme (not
    just the first) is rebranded. The first master wins on duplicate hexes."""
    out: dict[str, str] = {}
    for st in src_themes:
        for slot in _THEME_SLOTS:
            if st.get(slot) and dest_theme.get(slot):
                out.setdefault(st[slot], dest_theme[slot])
    return out


def _source_theme_font_map(src_fonts_list: list[dict],
                           dest_fonts: dict) -> dict[str, str]:
    """Map every source master theme's major/minor typeface to the brand's."""
    out: dict[str, str] = {}
    for sf in src_fonts_list:
        for kind in ("major", "minor"):
            if sf.get(kind) and dest_fonts.get(kind):
                out.setdefault(sf[kind], dest_fonts[kind])
    return out


def _hex_hls(hex6: str) -> tuple[float, float, float]:
    r = int(hex6[0:2], 16) / 255
    g = int(hex6[2:4], 16) / 255
    b = int(hex6[4:6], 16) / 255
    return colorsys.rgb_to_hls(r, g, b)


def _hls_hex(h: float, l: float, s: float) -> str:
    r, g, b = colorsys.hls_to_rgb(h, l, s)
    return "%02X%02X%02X" % (round(r * 255), round(g * 255), round(b * 255))


def color_clusters(src_prs, protected: set[str]) -> list[dict]:
    """Group the deck's hard-coded chromatic colors into hue families.
    Returns clusters sorted by usage weight (dominant first), each with a
    stable id: {"id", "hue_deg", "weight", "colors": {hex: count},
    "dominant": bool}. Neutrals and brand colors are excluded.
    """
    counts: Counter = Counter()
    for slide in src_prs.slides:
        for el in slide._element.iter(f"{{{A_NS}}}srgbClr"):
            val = (el.get("val") or "").upper()
            if len(val) == 6:
                counts[val] += 1

    chroma: dict[str, tuple[float, float, float, int]] = {}
    for val, n in counts.items():
        if val in protected:
            continue  # already a brand color
        h, l, s = _hex_hls(val)
        if s < 0.15 or l < 0.06 or l > 0.98:
            continue  # neutral / near-black / near-white
        chroma[val] = (h, l, s, n)
    if not chroma:
        return []

    buckets: dict[int, dict] = {}  # 12 circular buckets of 30 degrees
    for val, (h, l, s, n) in chroma.items():
        b = int(h * 12) % 12
        entry = buckets.setdefault(b, {"weight": 0, "colors": {}})
        entry["weight"] += n
        entry["colors"][val] = n
    clusters = [
        {"id": f"colors-hue-{b * 30}", "hue_deg": b * 30 + 15,
         "weight": e["weight"],
         "colors": dict(sorted(e["colors"].items(),
                               key=lambda kv: -kv[1]))}
        for b, e in buckets.items()]
    clusters.sort(key=lambda c: -c["weight"])
    for i, c in enumerate(clusters):
        c["dominant"] = i == 0
    return clusters


def _harmonization_map(clusters: list[dict], brand_hex: str,
                       rebrand_ids: set[str]) -> dict[str, str]:
    """Map every color of the selected clusters onto the brand primary.
    Colors whose lightness is close to the primary's snap to the EXACT
    primary (accent shapes should be truly on-brand, not merely brand-hued);
    lighter/darker variants keep their lightness so tints stay tints."""
    brand_h, brand_l, brand_s = _hex_hls(brand_hex)
    out: dict[str, str] = {}
    for cluster in clusters:
        if cluster["id"] not in rebrand_ids:
            continue
        for val in cluster["colors"]:
            h, l, s = _hex_hls(val)
            if abs(l - brand_l) <= 0.12:
                out[val] = brand_hex.upper()
            else:
                out[val] = _hls_hex(brand_h, l, max(s, brand_s * 0.5))
    return out


def collect_hardcoded_fonts(src_prs, exclude: set[str]) -> dict[str, int]:
    """Distinct literal typefaces used in slide runs (minus theme fonts)."""
    counts: Counter = Counter()
    for slide in src_prs.slides:
        for el in slide._element.iter(f"{{{A_NS}}}latin"):
            face = el.get("typeface")
            if face and not face.startswith("+") and face not in exclude:
                counts[face] += 1
    return dict(counts)


def _rewrite_rids(blob: bytes, rid_map: dict[str, str]) -> bytes:
    """Rewrite r:embed/r:id/r:link/... attribute values per rid_map."""
    root = etree.fromstring(blob)
    changed = False
    for el in root.iter():
        for name, val in el.attrib.items():
            if name.startswith(f"{{{R_NS}}}") and val in rid_map:
                el.set(name, rid_map[val])
                changed = True
    if not changed:
        return blob
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                          standalone=True)


def _restyle_tree(root, color_map: dict[str, str],
                  font_map: dict[str, str]) -> bool:
    """Remap hard-coded formatting that matches the SOURCE theme to the
    brandkit's theme, in place: srgbClr values (shape/text/table/chart
    colors) and latin/ea/cs typefaces (fonts). Positions are never touched."""
    changed = False
    for el in root.iter(f"{{{A_NS}}}srgbClr"):
        val = (el.get("val") or "").upper()
        if val in color_map and color_map[val] != val:
            el.set("val", color_map[val])
            changed = True
    if font_map:
        for tag in ("latin", "ea", "cs"):
            for el in root.iter(f"{{{A_NS}}}{tag}"):
                face = el.get("typeface")
                if face in font_map and font_map[face] != face:
                    el.set("typeface", font_map[face])
                    changed = True
    return changed


def _recolor_srgb(blob: bytes, color_map: dict[str, str]) -> bytes:
    """Blob-level restyle (colors only) — kept for chart parts and tests."""
    return _restyle_blob(blob, color_map, {})


def _restyle_blob(blob: bytes, color_map: dict[str, str],
                  font_map: dict[str, str]) -> bytes:
    root = etree.fromstring(blob)
    if not _restyle_tree(root, color_map, font_map):
        return blob
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                          standalone=True)


# scheme-color value -> theme slot. bg/tx follow the default color map
# (bg1->lt1, tx1->dk1, bg2->lt2, tx2->dk2); accents/dk/lt map to themselves.
_SCHEME_TO_SLOT = {
    "dk1": "dk1", "lt1": "lt1", "dk2": "dk2", "lt2": "lt2",
    "tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2",
    **{f"accent{i}": f"accent{i}" for i in range(1, 7)},
}


def _freeze_theme(root, theme_colors: dict, theme_fonts: dict) -> None:
    """Bake a slide's theme-linked formatting to fixed source values so it can
    no longer re-resolve against whatever theme the slide lands under.

    schemeClr -> srgbClr(source hex), keeping any tint/shade/lum modifier
    children; theme font refs (+mj-*/+mn-*) -> the source theme's actual
    typefaces. Used for "fully retain original look" protected slides."""
    for el in list(root.iter(f"{{{A_NS}}}schemeClr")):
        hexv = theme_colors.get(_SCHEME_TO_SLOT.get(el.get("val")))
        if not hexv:
            continue  # phClr / hlink / unknown: leave for the theme to resolve
        el.tag = f"{{{A_NS}}}srgbClr"      # keeps children (lumMod/tint/shade…)
        el.set("val", hexv)
    major, minor = theme_fonts.get("major"), theme_fonts.get("minor")
    for tag in ("latin", "ea", "cs"):
        for el in root.iter(f"{{{A_NS}}}{tag}"):
            face = el.get("typeface") or ""
            if face.startswith("+mj-") and major:
                el.set("typeface", major)
            elif face.startswith("+mn-") and minor:
                el.set("typeface", minor)


def _parse_protection(choices: dict) -> tuple[set[int], str, set[str]]:
    """Read the per-slide protection controls from the choices map.

    __protect__       = "2,4,7"  (1-based slide numbers to leave alone)
    __protect_mode__  = "full" | "partial"
    __protect_apply__ = "color,font,title,heading,resize"  (partial: still-on)
    Returns (protected 1-based slide numbers, mode, applied-adjustment set)."""
    raw = (choices.get("__protect__") or "").strip()
    protected = {int(x) for x in raw.split(",") if x.strip().isdigit()}
    mode = choices.get("__protect_mode__", "full")
    apply_set = {a for a in (choices.get("__protect_apply__") or "").split(",")
                 if a}
    return protected, mode, apply_set


_TITLE_PH_TYPES = {"title", "ctrTitle", "subTitle"}
HEADING_MIN_CENTIPT = 2000  # runs >= 20pt in plain text boxes count as headings


def deck_has_textbox_headings(src_prs) -> bool:
    for slide in src_prs.slides:
        for sp in slide._element.iter(f"{{{P_NS}}}sp"):
            if sp.find(f".//{{{P_NS}}}ph") is not None:
                continue
            for rpr in sp.iter(f"{{{A_NS}}}rPr"):
                sz = rpr.get("sz")
                if sz and int(sz) >= HEADING_MIN_CENTIPT:
                    return True
    return False


def _set_run_color(rpr, hex_color: str) -> None:
    solid = rpr.find(f"{{{A_NS}}}solidFill")
    if solid is not None:
        srgb = solid.find(f"{{{A_NS}}}srgbClr")
        if srgb is not None:
            srgb.set("val", hex_color)
            return
        for child in list(solid):
            solid.remove(child)
    else:
        solid = etree.SubElement(rpr, f"{{{A_NS}}}solidFill")
        # fills must precede a:latin/a:ea/a:cs in CT_TextCharacterProperties
        for tag in ("latin", "ea", "cs", "sym"):
            el = rpr.find(f"{{{A_NS}}}{tag}")
            if el is not None:
                el.addprevious(solid)
                break
    etree.SubElement(solid, f"{{{A_NS}}}srgbClr").set("val", hex_color)


def _heading_boxes(sld_el) -> list:
    """Non-placeholder shapes containing heading-sized runs, with their max
    run size and vertical position (for picking the primary heading)."""
    out = []
    for sp in sld_el.iter(f"{{{P_NS}}}sp"):
        if sp.find(f".//{{{P_NS}}}ph") is not None:
            continue
        max_sz = 0
        for rpr in sp.iter(f"{{{A_NS}}}rPr"):
            sz = rpr.get("sz")
            if sz:
                max_sz = max(max_sz, int(sz))
        if max_sz >= HEADING_MIN_CENTIPT:
            off = sp.find(f".//{{{A_NS}}}xfrm/{{{A_NS}}}off")
            y = int(off.get("y")) if off is not None else 1 << 60
            out.append((sp, max_sz, y))
    return out


def _slide_has_used_title_ph(sld_el) -> bool:
    for sp in sld_el.iter(f"{{{P_NS}}}sp"):
        ph = sp.find(f".//{{{P_NS}}}ph")
        if ph is not None and ph.get("type") in ("title", "ctrTitle"):
            text = "".join(t.text or "" for t in sp.iter(f"{{{A_NS}}}t"))
            if text.strip():
                return True
    return False


def _move_shape(sp, geom: tuple[int, int, int, int]) -> None:
    """Set a shape's transform to (x, y, cx, cy) EMU."""
    sp_pr = sp.find(f"{{{P_NS}}}spPr")
    if sp_pr is None:
        return
    xfrm = sp_pr.find(f"{{{A_NS}}}xfrm")
    if xfrm is None:
        xfrm = etree.Element(f"{{{A_NS}}}xfrm")
        sp_pr.insert(0, xfrm)  # a:xfrm must be the first child of spPr
    for tag in ("off", "ext"):
        el = xfrm.find(f"{{{A_NS}}}{tag}")
        if el is None:
            el = etree.SubElement(xfrm, f"{{{A_NS}}}{tag}")
    off = xfrm.find(f"{{{A_NS}}}off")
    ext = xfrm.find(f"{{{A_NS}}}ext")
    off.set("x", str(geom[0])), off.set("y", str(geom[1]))
    ext.set("cx", str(geom[2])), ext.set("cy", str(geom[3]))


def _restyle_heading_runs(sp, title_fmt: dict, heading_font: str | None,
                          apply_color: bool = True,
                          keep_fonts: frozenset = frozenset()):
    size_pt = title_fmt.get("size_pt")
    color = title_fmt.get("color") if apply_color else None
    for rpr in sp.iter(f"{{{A_NS}}}rPr"):
        sz = rpr.get("sz")
        if not sz or int(sz) < HEADING_MIN_CENTIPT:
            continue
        if size_pt:
            rpr.set("sz", str(round(size_pt * 100)))
        if title_fmt.get("bold"):
            rpr.set("b", "1")
        latin = rpr.find(f"{{{A_NS}}}latin")
        # a font the user chose to KEEP must win over the brand heading font,
        # otherwise the font decision would have no visible effect on headings
        kept = latin is not None and latin.get("typeface") in keep_fonts
        if heading_font and not kept:
            if latin is not None:
                latin.set("typeface", heading_font)
            else:
                etree.SubElement(rpr, f"{{{A_NS}}}latin").set(
                    "typeface", heading_font)
        if color:
            _set_run_color(rpr, color)


def _is_top_level_sp(sp) -> bool:
    """True when the shape is a direct child of the slide shape tree. Shapes
    inside a group use the group's child coordinate space, so their a:off is
    NOT slide-absolute and must never be overwritten with slide EMU."""
    parent = sp.getparent()
    return parent is not None and etree.QName(parent).localname == "spTree"


def _shape_rect(sp) -> tuple[int, int, int, int] | None:
    xfrm = sp.find(f".//{{{A_NS}}}xfrm")
    if xfrm is None:
        return None
    off = xfrm.find(f"{{{A_NS}}}off")
    ext = xfrm.find(f"{{{A_NS}}}ext")
    if off is None or ext is None:
        return None
    return (int(off.get("x")), int(off.get("y")),
            int(ext.get("cx")), int(ext.get("cy")))


def _rects_intersect(a, b) -> bool:
    return not (a[0] + a[2] <= b[0] or b[0] + b[2] <= a[0] or
                a[1] + a[3] <= b[1] or b[1] + b[3] <= a[1])


def _sp_fill_is_colored(sp) -> bool:
    """True when the shape carries its own solid fill that reads as a
    colored surface (text sitting on it depends on that fill for contrast).
    Near-white fills behave like the slide background and don't count."""
    sp_pr = sp.find(f"{{{P_NS}}}spPr")
    if sp_pr is None:
        return False
    fill = sp_pr.find(f"{{{A_NS}}}solidFill")
    if fill is None:
        return False
    srgb = fill.find(f"{{{A_NS}}}srgbClr")
    if srgb is not None:
        val = (srgb.get("val") or "").upper()
        if len(val) != 6:
            return False
        _, l, _ = _hex_hls(val)
        return l < 0.92
    scheme = fill.find(f"{{{A_NS}}}schemeClr")
    if scheme is not None:
        return scheme.get("val") not in ("bg1", "lt1")
    return False


def _colored_shape_rects(sld_el) -> list[tuple]:
    return [(sp, rect) for sp in sld_el.iter(f"{{{P_NS}}}sp")
            if _sp_fill_is_colored(sp) and (rect := _shape_rect(sp))]


def _on_colored_shape(sp, colored_rects) -> bool:
    """Text lives inside a colored shape, or in a box overlapping one."""
    if _sp_fill_is_colored(sp):
        return True
    rect = _shape_rect(sp)
    if rect is None:
        return False
    return any(other is not sp and _rects_intersect(rect, r)
               for other, r in colored_rects)


def _restyle_all_runs(sp, fmt: dict, font: str | None) -> None:
    """Apply a format to every run of a shape (used for secondary text
    moved into the brandkit subtitle slot)."""
    for rpr in sp.iter(f"{{{A_NS}}}rPr"):
        if fmt.get("size_pt"):
            rpr.set("sz", str(round(fmt["size_pt"] * 100)))
        if font:
            latin = rpr.find(f"{{{A_NS}}}latin")
            if latin is not None:
                latin.set("typeface", font)
            else:
                etree.SubElement(rpr, f"{{{A_NS}}}latin").set("typeface", font)
        if fmt.get("color"):
            _set_run_color(rpr, fmt["color"])


def _apply_heading_style(sld_el, title_fmt: dict, heading_font: str | None,
                         title_geom: tuple[int, int, int, int] | None = None,
                         sub_geom: tuple[int, int, int, int] | None = None,
                         sub_fmt: dict | None = None,
                         body_font: str | None = None,
                         keep_fonts: frozenset = frozenset()):
    """Non-placeholder text runs at heading sizes adopt the brandkit title
    style of the slide's target layout ("restyle in place": font, size and
    color only — geometry untouched).

    With `title_geom` ("convert"): the slide's primary heading box (largest,
    then topmost) is MOVED to the brandkit title position, and any other
    text box sitting in that title zone (kickers, secondary lines) is moved
    to the brandkit SUBTITLE position with subtitle styling — unless the
    slide already uses a real title placeholder.

    Text sitting on a colored shape (inside a filled shape, or in a box
    overlapping one) stays put — it belongs to that shape visually — and
    keeps its own run colors (white-on-color stays white; the shape's fill
    is governed by the color-family decisions instead)."""
    boxes = _heading_boxes(sld_el)
    if not boxes:
        return
    colored_rects = _colored_shape_rects(sld_el)
    for sp, _, _ in boxes:
        _restyle_heading_runs(
            sp, title_fmt, heading_font,
            apply_color=not _on_colored_shape(sp, colored_rects),
            keep_fonts=keep_fonts)
    if not title_geom or _slide_has_used_title_ph(sld_el):
        return

    # only shapes whose a:off is slide-absolute (top-level, not grouped) can
    # be relocated to a brandkit slot; a grouped shape's coords are in its
    # group's child space and moving it would fling it off-slide
    movable = [b for b in boxes
               if not _on_colored_shape(b[0], colored_rects)
               and _is_top_level_sp(b[0])]
    if not movable:
        return
    primary = sorted(movable, key=lambda b: (-b[1], b[2]))[0][0]
    # find secondary text boxes currently occupying the title target zone,
    # BEFORE anything moves
    secondaries = []
    for sp in sld_el.iter(f"{{{P_NS}}}sp"):
        if sp is primary or sp.find(f".//{{{P_NS}}}ph") is not None:
            continue
        if not _is_top_level_sp(sp):
            continue  # grouped: coords not slide-absolute, cannot relocate
        if sp.find(f"{{{P_NS}}}txBody") is None:
            continue
        if _on_colored_shape(sp, colored_rects):
            continue  # stays with its shape
        text = "".join(t.text or "" for t in sp.iter(f"{{{A_NS}}}t")).strip()
        if not text:
            continue
        rect = _shape_rect(sp)
        if rect and _rects_intersect(rect, title_geom):
            secondaries.append(sp)

    _move_shape(primary, title_geom)
    if sub_geom:
        for sp in secondaries:
            _move_shape(sp, sub_geom)
            _restyle_all_runs(sp, sub_fmt or {}, body_font)


def _adopt_brand_title_sizes(sld_el) -> bool:
    """Strip explicit font sizes (and autofit shrink factors) from title-like
    placeholders so the brandkit layout's own sizes apply. Body text keeps
    its sizes — authors often sized it deliberately to fit."""
    changed = False
    for sp in sld_el.iter(f"{{{P_NS}}}sp"):
        ph = sp.find(f".//{{{P_NS}}}ph")
        if ph is None or ph.get("type") not in _TITLE_PH_TYPES:
            continue
        for el in sp.iter():
            local = etree.QName(el).localname
            if local in ("rPr", "endParaRPr", "defRPr") and el.get("sz"):
                del el.attrib["sz"]
                changed = True
            elif local == "normAutofit":
                if el.attrib.pop("fontScale", None) is not None:
                    changed = True
                el.attrib.pop("lnSpcReduction", None)
    return changed


class _PartCloner:
    """Clones parts (and their relationship subgraphs) from one package
    into another, deduplicating identical binaries."""

    def __init__(self, dest_package):
        self.pkg = dest_package
        self._memo: dict[str, Part] = {}      # src partname -> dest part
        self._by_hash: dict[str, Part] = {}   # binary dedupe
        self.cloned: list[Part] = []
        # Persistent set of taken partnames + a per-(prefix, ext) next-index
        # cursor, so allocating N names is O(N) instead of rescanning every
        # package part (and restarting the search from 1) on each call. The
        # cloner is the only thing that mints media/chart part names here, so
        # seeding once from the package and reserving as we go stays correct.
        self._used = {str(p.partname) for p in dest_package.iter_parts()}
        self._next_n: dict[tuple[str, str], int] = {}

    def _next_partname(self, src_partname: str) -> PackURI:
        m = re.match(r"^(.*?)(\d*)(\.\w+)$", src_partname)
        prefix, ext = m.group(1), m.group(3)
        n = self._next_n.get((prefix, ext), 1)
        while f"{prefix}{n}{ext}" in self._used:
            n += 1
        name = f"{prefix}{n}{ext}"
        self._used.add(name)                  # reserve before the part exists
        self._next_n[(prefix, ext)] = n + 1
        return PackURI(name)

    def clone(self, src_part) -> Part:
        key = str(src_part.partname)
        if key in self._memo:
            return self._memo[key]
        blob = src_part.blob
        is_binary = not src_part.content_type.endswith("xml")
        if is_binary:
            digest = hashlib.sha1(blob).hexdigest()
            if digest in self._by_hash:
                self._memo[key] = self._by_hash[digest]
                return self._by_hash[digest]

        new_part = Part(self._next_partname(key), src_part.content_type,
                        self.pkg, blob)
        self._memo[key] = new_part
        self.cloned.append(new_part)
        if is_binary:
            self._by_hash[digest] = new_part

        rid_map: dict[str, str] = {}
        for r_id, rel in src_part.rels.items():
            if rel.is_external:
                new_rid = new_part.relate_to(rel.target_ref, rel.reltype,
                                             is_external=True)
            else:
                child = self.clone(rel.target_part)
                new_rid = new_part.relate_to(child, rel.reltype)
            rid_map[r_id] = new_rid
        if rid_map and not is_binary:
            new_part._blob = _rewrite_rids(new_part.blob, rid_map)
        return new_part


def _scale_shape_tree(spTree, factor: float) -> None:
    """Uniformly scale the position and size of every top-level shape by
    `factor`, to fit transplanted content onto a differently-sized brand
    canvas. Grouped children scale automatically with their group's extent,
    so only direct children of the shape tree are touched."""
    for child in spTree:
        local = etree.QName(child).localname
        if local == "graphicFrame":
            xfrm = child.find(f"{{{P_NS}}}xfrm")
        elif local == "grpSp":
            xfrm = child.find(f"{{{P_NS}}}grpSpPr/{{{A_NS}}}xfrm")
        elif local in ("sp", "pic", "cxnSp"):
            xfrm = child.find(f"{{{P_NS}}}spPr/{{{A_NS}}}xfrm")
        else:
            continue
        if xfrm is None:
            continue
        off = xfrm.find(f"{{{A_NS}}}off")
        ext = xfrm.find(f"{{{A_NS}}}ext")
        if off is not None and off.get("x") and off.get("y"):
            off.set("x", str(round(int(off.get("x")) * factor)))
            off.set("y", str(round(int(off.get("y")) * factor)))
        if ext is not None and ext.get("cx") and ext.get("cy"):
            ext.set("cx", str(round(int(ext.get("cx")) * factor)))
            ext.set("cy", str(round(int(ext.get("cy")) * factor)))


def _table_styles_part(prs):
    """The presentation's ppt/tableStyles.xml part, or None."""
    for rel in prs.part.rels.values():
        if rel.reltype.endswith("tableStyles"):
            return rel.target_part
    return None


def _merge_table_styles(src, dest, color_map: dict, font_map: dict,
                        warnings: list) -> None:
    """Copy any table-style definitions the content deck's tables reference
    into the brandkit's tableStyles.xml. A transplanted table stores only a
    style GUID; if that GUID isn't defined in the destination package the
    table renders unstyled. Merged styles are restyled to the brand (their
    theme-bound schemeClr colors already re-resolve against the brand theme;
    hard-coded source-theme colors/fonts are rewritten here)."""
    src_part, dest_part = _table_styles_part(src), _table_styles_part(dest)
    if src_part is None or dest_part is None:
        return
    src_root = etree.fromstring(src_part.blob)
    dest_root = etree.fromstring(dest_part.blob)
    have = {ts.get("styleId")
            for ts in dest_root.findall(f"{{{A_NS}}}tblStyle")}
    added = 0
    for ts in src_root.findall(f"{{{A_NS}}}tblStyle"):
        if ts.get("styleId") in have:
            continue
        node = copy.deepcopy(ts)
        _restyle_tree(node, color_map, font_map)
        dest_root.append(node)
        added += 1
    if added:
        dest_part._blob = etree.tostring(
            dest_root, xml_declaration=True, encoding="UTF-8", standalone=True)
        warnings.append(f"{added} table style(s) carried over from the "
                        f"content deck so its tables keep their structure")


def _transplant_notes(src_slide, new_slide, dest, cloner) -> bool:
    """Clone the source slide's notesSlide part into the dest package and wire
    it to the new slide and the dest notes master, so speaker notes survive
    the retheme. Returns True when notes were carried over."""
    if not src_slide.has_notes_slide:
        return False
    src_notes = src_slide.notes_slide.part
    new_notes = Part(
        cloner._next_partname("/ppt/notesSlides/notesSlide1.xml"),
        src_notes.content_type, dest.part.package, src_notes.blob)
    rid_map: dict[str, str] = {}
    for r_id, rel in src_notes.rels.items():
        if rel.reltype == RT.NOTES_MASTER:
            # reuse the dest's own notes master rather than cloning a duplicate
            rid_map[r_id] = new_notes.relate_to(
                dest.notes_master.part, RT.NOTES_MASTER)
        elif rel.reltype == RT.SLIDE:
            rid_map[r_id] = new_notes.relate_to(new_slide.part, RT.SLIDE)
        elif rel.is_external:
            rid_map[r_id] = new_notes.relate_to(
                rel.target_ref, rel.reltype, is_external=True)
        else:
            child = cloner.clone(rel.target_part)
            rid_map[r_id] = new_notes.relate_to(child, rel.reltype)
    if rid_map:
        new_notes._blob = _rewrite_rids(new_notes.blob, rid_map)
    new_slide.part.relate_to(new_notes, RT.NOTES_SLIDE)
    return True


def _pick_layout(src_slide, layouts_by_name: dict, manifest: dict,
                 warnings: list[str], index: int,
                 overrides: dict[str, str] | None = None):
    lname = src_slide.slide_layout.name
    if overrides and lname in overrides:
        layout = layouts_by_name.get(overrides[lname])
        if layout is not None:
            return layout
    layout = layouts_by_name.get(lname)
    if layout is not None:
        return layout
    purpose = infer_purpose(src_slide.slide_layout.element.get("type"), lname)
    tag = _PURPOSE_TO_TAG.get(purpose, "content")
    try:
        ldef = _find_layout(manifest, tag)
        layout = layouts_by_name.get(ldef["name"])
    except LayoutNotFoundError:
        layout = None
    if layout is None:
        layout = next(iter(layouts_by_name.values()))
        warnings.append(f'slide {index}: no brandkit layout matches '
                        f'"{lname}" — placed on "{layout.name}"')
    return layout


def retheme_pptx(package_dir: str | Path, deck_path: str | Path,
                 out_path: str | Path,
                 choices: dict[str, str] | None = None) -> list[str]:
    """Copy every slide of `deck_path` into the brandkit's source template.

    `choices` optionally resolves the user-confirmable decisions produced by
    `conflicts.analyze_conflicts` (decision id -> chosen value); without it,
    sensible defaults apply (rebrand the dominant color family, map foreign
    fonts to the brand body font, adopt brand title sizes).

    Returns warnings. Raises FileNotFoundError if the package has no
    embedded source template (callers may fall back to the rebuild path)."""
    choices = choices or {}
    from ..extractor.extract import _as_presentation

    package_dir, deck_path = Path(package_dir), Path(deck_path)
    manifest = json.loads(
        (package_dir / "template.json").read_text(encoding="utf-8"))
    source_rel = manifest.get("source_template")
    source_path = package_dir / source_rel if source_rel else None
    if not source_path or not source_path.is_file():
        raise FileNotFoundError("package has no embedded source template")

    warnings: list[str] = []
    tmp = Path(tempfile.mkdtemp(prefix="retheme-"))
    try:
        dest = Presentation(str(_as_presentation(source_path, tmp)))
        src = Presentation(str(_as_presentation(deck_path, tmp)))

        # slide-size mismatch: keep positions (default, transparent) or, when
        # the user opts in with resize="fit", uniformly scale content to fit
        # the brand canvas (uniform factor so nothing is stretched/distorted)
        size_differs = (src.slide_width, src.slide_height) != \
            (dest.slide_width, dest.slide_height)
        scale_factor = None
        if size_differs:
            if choices.get("resize", "keep") == "fit" \
                    and src.slide_width and src.slide_height:
                scale_factor = min(dest.slide_width / src.slide_width,
                                   dest.slide_height / src.slide_height)
                warnings.append(
                    "content deck and brandkit have different slide sizes "
                    f"({src.slide_width}x{src.slide_height} vs "
                    f"{dest.slide_width}x{dest.slide_height} EMU): content "
                    f"scaled by {scale_factor:.3f} to fit the brand canvas")
            else:
                warnings.append(
                    "content deck and brandkit have different slide sizes "
                    f"({src.slide_width}x{src.slide_height} vs "
                    f"{dest.slide_width}x{dest.slide_height} EMU): retained "
                    "content keeps its absolute positions and may need "
                    "adjustment in PowerPoint")

        layouts_by_name = {}
        for master in dest.slide_masters:
            for layout in master.slide_layouts:
                layouts_by_name.setdefault(layout.name, layout)

        # source-theme formatting -> brandkit formatting maps: hard-coded
        # colors/fonts matching the source theme get rewritten to the brand.
        # Aggregated across ALL source masters so multi-master decks rebrand
        # colors bound to any master's theme, not just the first.
        src_themes = _theme_colors_all(src)
        dest_theme = _theme_colors_of(dest)
        theme_color_map = _source_theme_color_map(src_themes, dest_theme)
        color_map = dict(theme_color_map)
        src_fonts_list = _theme_fonts_all(src)
        dest_fonts = _theme_fonts_of(dest)
        font_map = _source_theme_font_map(src_fonts_list, dest_fonts)

        # harmonize hard-coded off-theme colors. Default: only the dominant
        # color family is rebranded; user choices can rebrand/keep any family
        primary = next((c["hex"] for c in manifest.get("colors", [])
                        if c.get("role") == "primary"), None) \
            or dest_theme.get("accent1")
        if primary:
            protected = set(dest_theme.values()) | set(color_map.values())
            clusters = color_clusters(src, protected)
            rebrand_ids = {
                c["id"] for c in clusters
                if choices.get(c["id"], "rebrand" if c["dominant"] else
                               "keep") == "rebrand"}
            harmony = _harmonization_map(clusters, primary, rebrand_ids)
            if harmony:
                warnings.append(
                    f"{len(harmony)} hard-coded color(s) recolored to the "
                    f"brand primary (#{primary})")
                # theme-derived mappings take precedence over harmonization
                color_map = {**harmony, **theme_color_map}

        # foreign hard-coded fonts: map to the brand body font unless the
        # user chose to keep them. Kept faces are also protected from the
        # heading restyle below, so the choice actually shows on headings.
        exclude = {f for sf in src_fonts_list for f in sf.values()} \
            | set(dest_fonts.values())
        keep_fonts: set[str] = set()
        for face in collect_hardcoded_fonts(src, exclude):
            if choices.get(f"font:{face}", "map") == "map" \
                    and dest_fonts.get("minor"):
                font_map[face] = dest_fonts["minor"]
            else:
                keep_fonts.add(face)

        layout_overrides = {
            key[len("layout:"):]: value for key, value in choices.items()
            if key.startswith("layout:")}
        adopt_title_sizes = choices.get("title_sizes", "adopt") == "adopt"
        # heading_style: "convert" (restyle + move to title position),
        # "adopt" (restyle in place), or "keep"
        heading_mode = choices.get("heading_style", "convert")
        manifest_layouts = {l["name"]: l for l in manifest.get("layouts", [])}
        heading_font = dest_fonts.get("major") or dest_fonts.get("minor")

        def _title_ph_of(layout_rec):
            for ph in layout_rec.get("placeholders", []):
                if ph["type"] in ("title", "ctrTitle") and \
                        ph.get("format_levels", {}).get("1"):
                    return ph
            return None

        def _sub_ph_of(layout_rec):
            for ph in layout_rec.get("placeholders", []):
                if ph["type"] == "subTitle":
                    return ph
            return None

        def _geom_of(ph) -> tuple[int, int, int, int] | None:
            pos, size = ph.get("position"), ph.get("size")
            if not pos or not size:
                return None
            return (round(pos["x_in"] * 914400), round(pos["y_in"] * 914400),
                    round(size["width_in"] * 914400),
                    round(size["height_in"] * 914400))

        # brandkit-wide title/subtitle styles, for slides landing on layouts
        # that have no such placeholder of their own (e.g. Blank)
        default_title_ph = next(
            (ph for l in manifest.get("layouts", [])
             if (ph := _title_ph_of(l))), None)
        default_sub_ph = next(
            (ph for l in manifest.get("layouts", [])
             if (ph := _sub_ph_of(l))), None)
        default_sub_fmt = (default_sub_ph or {}).get(
            "format_levels", {}).get("1", {"size_pt": 14})

        # per-slide protection: listed slides opt out of the choices. "full"
        # freezes them to their original look (theme colors/fonts baked in);
        # "partial" still applies only the chosen adjustment categories.
        protected, protect_mode, protect_apply = _parse_protection(choices)
        src_theme_colors = _theme_colors_of(src)
        src_theme_fonts = _theme_fonts_of(src)
        protected_parts: set = set()   # cloned parts of fully-frozen slides
        _ALL_ADJ = {"color", "font", "title", "heading", "resize"}

        _delete_all_slides(dest)
        cloner = _PartCloner(dest.part.package)
        # map each source slide part to its 0-based position, so an internal
        # slide-to-slide hyperlink can be repointed at the CORRESPONDING new
        # slide instead of deep-cloning the target slide into an orphan part
        src_slide_index = {sl.part: idx
                           for idx, sl in enumerate(src.slides)}
        records = []          # (new_slide, new_sld, rid_map) per output slide
        pending_links = []    # (from_idx, r_id, target_idx | None)
        notes_kept = 0

        for i, s in enumerate(src.slides, start=1):
            layout = _pick_layout(s, layouts_by_name, manifest, warnings, i,
                                  layout_overrides)
            new_slide = dest.slides.add_slide(layout)

            # which adjustments this slide is open to (all, a chosen subset,
            # or none when fully frozen)
            is_full = i in protected and protect_mode == "full"
            if i in protected:
                allow = set() if is_full else set(protect_apply)
            else:
                allow = _ALL_ADJ

            # transplant the whole shape tree (and any slide bg override)
            new_sld = new_slide._element
            src_cSld = s._element.find(f"{{{P_NS}}}cSld")
            new_cSld = new_sld.find(f"{{{P_NS}}}cSld")
            new_spTree = new_cSld.find(f"{{{P_NS}}}spTree")
            new_cSld.remove(new_spTree)
            new_cSld.append(copy.deepcopy(
                src_cSld.find(f"{{{P_NS}}}spTree")))
            if scale_factor is not None and "resize" in allow:
                _scale_shape_tree(
                    new_cSld.find(f"{{{P_NS}}}spTree"), scale_factor)
            src_bg = src_cSld.find(f"{{{P_NS}}}bg")
            if src_bg is not None:
                new_cSld.insert(0, copy.deepcopy(src_bg))
                warnings.append(f"slide {i}: keeps its own background "
                                f"override from the content deck")

            # bring over the slide's related parts (images, charts, ...).
            # Internal slide jumps are deferred: their target new slide may
            # not exist yet, and cloning it would duplicate a slide part.
            rid_map: dict[str, str] = {}
            for r_id, rel in s.part.rels.items():
                if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
                    continue
                if rel.reltype == RT.SLIDE and not rel.is_external:
                    pending_links.append(
                        (i - 1, r_id, src_slide_index.get(rel.target_part)))
                    continue
                if rel.is_external:
                    new_rid = new_slide.part.relate_to(
                        rel.target_ref, rel.reltype, is_external=True)
                else:
                    child = cloner.clone(rel.target_part)
                    new_rid = new_slide.part.relate_to(child, rel.reltype)
                    if is_full:
                        # keep this slide's charts/diagrams unrestyled too
                        protected_parts.add(child)
                rid_map[r_id] = new_rid

            # carry over speaker notes (their own part + relationships)
            if _transplant_notes(s, new_slide, dest, cloner):
                notes_kept += 1

            # restyle the transplanted content in place: hard-coded colors
            # and fonts matching the source theme -> brandkit equivalents;
            # title placeholders adopt the brand layout's font sizes. (These
            # touch only the XML tree, not relationship ids, so they run now.)
            if is_full:
                # bake original theme colors/fonts in so the brand theme this
                # slide now lands under cannot shift its appearance
                _freeze_theme(new_sld, src_theme_colors, src_theme_fonts)
            _restyle_tree(new_sld,
                          color_map if "color" in allow else {},
                          font_map if "font" in allow else {})
            if adopt_title_sizes and "title" in allow:
                _adopt_brand_title_sizes(new_sld)
            if heading_mode in ("convert", "adopt") and "heading" in allow:
                layout_rec = manifest_layouts.get(layout.name, {})
                title_ph = _title_ph_of(layout_rec) or default_title_ph
                if title_ph:
                    title_fmt = title_ph.get("format_levels", {}).get("1", {})
                    geom = sub_geom = None
                    sub_fmt = default_sub_fmt
                    if heading_mode == "convert":
                        geom = _geom_of(title_ph)
                        sub_ph = _sub_ph_of(layout_rec)
                        if sub_ph and _geom_of(sub_ph):
                            sub_geom = _geom_of(sub_ph)
                            sub_fmt = sub_ph.get("format_levels", {}) \
                                .get("1", default_sub_fmt)
                        elif geom:
                            # no subtitle on this layout: synthesize a slot
                            # directly below the title
                            sub_geom = (geom[0], geom[1] + geom[3],
                                        geom[2], round(0.45 * 914400))
                    _apply_heading_style(
                        new_sld, title_fmt, heading_font, title_geom=geom,
                        sub_geom=sub_geom, sub_fmt=sub_fmt,
                        body_font=dest_fonts.get("minor"),
                        keep_fonts=frozenset(keep_fonts))

            records.append((new_slide, new_sld, rid_map))

        # every new slide now exists: resolve internal slide-to-slide links
        # to the matching output slide, then rewrite all relationship ids
        drop_rids: dict[int, set[str]] = {}
        for from_idx, r_id, target_idx in pending_links:
            from_part = records[from_idx][0].part
            if target_idx is None:
                drop_rids.setdefault(from_idx, set()).add(r_id)
                warnings.append("a slide-to-slide link targets a slide "
                                "outside the deck and was dropped")
                continue
            new_rid = from_part.relate_to(records[target_idx][0].part,
                                          RT.SLIDE)
            records[from_idx][2][r_id] = new_rid
        for idx, (_new_slide, new_sld, rid_map) in enumerate(records):
            dropped = drop_rids.get(idx, set())
            if not rid_map and not dropped:
                continue
            for el in new_sld.iter():
                for name, val in list(el.attrib.items()):
                    if not name.startswith(f"{{{R_NS}}}"):
                        continue
                    if val in rid_map:
                        el.set(name, rid_map[val])
                    elif val in dropped:
                        del el.attrib[name]  # avoid a dangling r:id
        if notes_kept:
            warnings.append(f"{notes_kept} slide(s) kept their speaker notes")

        # restyle cloned XML parts too (charts, chart styles, SmartArt):
        # colors AND fonts hard-coded to the source theme -> brandkit
        restyled = 0
        for part in cloner.cloned:
            if part in protected_parts:
                continue  # chart/diagram of a fully-frozen slide: leave as-is
            if part.content_type.endswith("xml"):
                new_blob = _restyle_blob(part.blob, color_map, font_map)
                if new_blob is not part.blob:
                    part._blob = new_blob
                    restyled += 1
        if restyled:
            warnings.append(f"{restyled} embedded part(s) (charts/diagrams) "
                            f"restyled to brandkit colors/fonts")

        # carry over any custom table-style definitions the deck's tables use
        _merge_table_styles(src, dest, color_map, font_map, warnings)

        dest.save(str(out_path))
        return warnings
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
