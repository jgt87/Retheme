"""Extract slide *content* (not styling) from an existing presentation into
the content-JSON format the application tool consumes.

This enables deck restyling: extract a template package from deck A (or a
corporate template), extract content from deck B, apply one to the other.

Carried over: titles, subtitles, body bullets with nesting, pictures
(as-is), tables (cell text grid) and chart markers (position + title) so
the HTML preview can show or indicate them. The PPTX export path
(`retheme.py`) retains everything natively and doesn't use these fields.
"""

from __future__ import annotations

import base64
import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ..ooxml import NS
from ..extractor.layouts import infer_purpose

# layout purpose (extractor vocabulary) -> content-JSON layout tag
_PURPOSE_TO_TAG = {
    "Title Slide": "title",
    "Section Header": "section_header",
    "Two Content": "two_content",
    "Comparison": "comparison",
    "Title Only": "title_only",
    "Blank": "blank",
    "Quote": "quote",
    "Closing Slide": "closing",
    "Agenda / Table of Contents": "agenda",
    "Picture Layout": "picture",
    "Picture with Caption": "picture",
}


def _paragraphs_to_bullets(text_frame) -> list:
    """Nested bullet list from a text frame, honoring paragraph levels."""
    root: list = []
    # stack[level] = list to append items of that level to
    stack = {0: root}
    for para in text_frame.paragraphs:
        text = "".join(run.text for run in para.runs).strip()
        if not text:
            continue
        level = min(para.level or 0, 2)
        parent = stack.get(level)
        if parent is None:  # level jumped without intermediate: clamp up
            level = max(l for l in stack if l <= level)
            parent = stack[level]
        item = {"text": text, "children": []}
        parent.append(item)
        stack[level + 1] = item["children"]
        for deeper in [l for l in stack if l > level + 1]:
            del stack[deeper]
    def strip_empty(items):
        for it in items:
            strip_empty(it["children"])
        return [it if it["children"] else it["text"] for it in items]
    return strip_empty(root)


def _ph_type(shape) -> str | None:
    ph = shape._element.find(".//p:nvSpPr/p:nvPr/p:ph", NS)
    if ph is None:
        ph = shape._element.find(".//p:nvPicPr/p:nvPr/p:ph", NS)
    return ph.get("type", "obj") if ph is not None else None


def extract_content(pptx_path: str | Path, package_dir: str | Path | None = None,
                    ) -> tuple[dict, list[str]]:
    """Returns (content_json_dict, warnings)."""
    from ..extractor.extract import _as_presentation
    import tempfile, shutil

    pptx_path = Path(pptx_path)
    known_layouts: set[str] = set()
    body_capable: set[str] = set()  # package layouts with a body placeholder
    if package_dir:
        manifest = json.loads(
            (Path(package_dir) / "template.json").read_text(encoding="utf-8"))
        known_layouts = {l["name"] for l in manifest["layouts"]}
        body_capable = {
            l["name"] for l in manifest["layouts"]
            if any(ph["type"] in ("body", "obj") for ph in l["placeholders"])}

    tmp = Path(tempfile.mkdtemp(prefix="pptx-content-"))
    warnings: list[str] = []
    try:
        prs = Presentation(str(_as_presentation(pptx_path, tmp)))
        sw, sh = prs.slide_width, prs.slide_height

        def geom(shape):
            if shape.left is None or shape.width is None:
                return None
            return {"x_pct": round(shape.left / sw, 4),
                    "y_pct": round(shape.top / sh, 4),
                    "w_pct": round(shape.width / sw, 4),
                    "h_pct": round(shape.height / sh, 4)}

        slides = []
        for i, slide in enumerate(prs.slides, start=1):
            layout = slide.slide_layout
            layout_name = layout.name
            layout_type = layout.element.get("type")
            purpose = infer_purpose(layout_type, layout_name)

            # exact layout name wins when the target package has it
            # (restyling a deck built on the same/renamed template);
            # otherwise fall back to a generic layout tag.
            tag = (layout_name if layout_name in known_layouts
                   else _PURPOSE_TO_TAG.get(purpose, "content"))

            entry: dict = {"layout": tag}
            bodies = []
            skipped = 0
            for shape in slide.shapes:
                ph_type = _ph_type(shape)
                # retained rich content: pictures as-is, tables as text
                # grids, charts as markers (fully retained in PPTX export)
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    g = geom(shape)
                    if g:
                        try:
                            img = shape.image
                            entry.setdefault("images", []).append({
                                **g, "mime": img.content_type,
                                "data_b64": base64.b64encode(img.blob).decode()})
                        except Exception:
                            skipped += 1
                    continue
                if getattr(shape, "has_table", False) and shape.has_table:
                    g = geom(shape)
                    if g:
                        rows = [[cell.text.strip() for cell in row.cells]
                                for row in shape.table.rows]
                        entry.setdefault("tables", []).append({**g, "rows": rows})
                    continue
                if getattr(shape, "has_chart", False) and shape.has_chart:
                    g = geom(shape)
                    if g:
                        title = ""
                        try:
                            if shape.chart.has_title:
                                title = shape.chart.chart_title.text_frame.text
                        except Exception:
                            pass
                        entry.setdefault("charts", []).append({**g, "title": title})
                    continue
                if shape.shape_type is not None and not shape.has_text_frame:
                    if ph_type is None:
                        continue  # decorative shape/picture: template's job
                    skipped += 1
                    continue
                if not shape.has_text_frame:
                    continue
                text = shape.text_frame.text.strip()
                if not text:
                    continue
                if ph_type in ("title", "ctrTitle"):
                    entry["title"] = text.replace("\n", " ")
                elif ph_type == "subTitle":
                    entry["subtitle"] = text.replace("\n", " ")
                elif ph_type in ("dt", "ftr", "sldNum"):
                    continue  # supplied by the template, not content
                else:
                    bullets = _paragraphs_to_bullets(shape.text_frame)
                    if bullets:
                        bodies.append((shape.top or 0, shape.left or 0, bullets))
            bodies.sort(key=lambda b: (b[0], b[1]))  # reading order
            bodies = [b[2] for b in bodies]
            if len(bodies) == 1:
                entry["bullets"] = bodies[0]
            elif len(bodies) == 2:
                entry["left"], entry["right"] = bodies[0], bodies[1]
                if tag == "content":
                    entry["layout"] = "two_content"
            elif len(bodies) > 2:
                # free-form slides: merge all text areas so nothing is lost
                entry["bullets"] = [item for body in bodies for item in body]
                warnings.append(
                    f"slide {i}: merged {len(bodies)} text areas into one "
                    f"list (layout '{layout_name}')")
            # exact-name match is pointless if that layout can't hold the text
            has_body = "bullets" in entry or "left" in entry
            if has_body and entry["layout"] in known_layouts \
                    and entry["layout"] not in body_capable:
                entry["layout"] = "two_content" if "left" in entry else "content"
                warnings.append(
                    f"slide {i}: layout '{layout_name}' has no body "
                    f"placeholder; using generic '{entry['layout']}' instead")
            elif has_body and entry["layout"] in ("title_only", "blank"):
                entry["layout"] = "content"
            if skipped:
                warnings.append(
                    f"slide {i}: {skipped} object(s) not representable in "
                    f"the HTML preview (they ARE retained in the PPTX export)")
            slides.append(entry)
        content = {"title": pptx_path.stem, "slides": slides}
        return content, warnings
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def extract_content_file(pptx_path: str | Path, out_path: str | Path,
                         package_dir: str | Path | None = None) -> list[str]:
    content, warnings = extract_content(pptx_path, package_dir)
    Path(out_path).write_text(
        json.dumps(content, indent=2, ensure_ascii=False), encoding="utf-8")
    return warnings
